import io
import json
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from backend.app import auth, excel_io, models, schemas, services
from backend.app.path_bom import EffectiveBOM, save_configuration
from backend.app.production_rebuild import transform_source


def make(db, code, kind='material', children=(), name=None):
    return services.create_item(db, schemas.ItemCreate(
        code=code, item_type=kind, name=name or code, status='active',
        source_type='purchased' if kind == 'material' else None, similarity_confirmed=True,
        components=[schemas.BOMComponentIn(child_item_id=c.id, quantity=Decimal(q)) for c, q in children]),
        db.scalar(select(models.User).where(models.User.role == 'admin')))


def configure(db, owner, members, selected=None):
    actor = db.scalar(select(models.User).where(models.User.role == 'admin'))
    row = EffectiveBOM(db).rows(owner.id)[0]
    return save_configuration(db, owner.id, {'line_path': row['line_path'], 'mode': 'custom',
        'selected_item_id': (selected or members[0][0]).id,
        'members': [{'item_id': item.id, 'market_share': share} for item, share in members], 'reason': '测试选配'}, actor)


def test_excel_only_candidates_columns_and_production_grouping(db):
    a = make(db, '10.7771.0', name='24V耐腐蚀蠕动泵')
    b = make(db, '10.7772.0'); c = make(db, '10.7773.0')
    s1 = make(db, '05.771.01', 'semi_finished', [(a, 2)])
    s2 = make(db, '05.772.01', 'semi_finished', [(a, 3)])
    s3 = make(db, '05.773.01', 'semi_finished', [(a, 4)])
    configure(db, s1, [(a, '50'), (b, '50')])
    configure(db, s2, [(a, '50'), (c, '50')])
    configure(db, s3, [(a, '50'), (b, '50')])
    root = make(db, '00.771.01', 'machine', [(s1, 1), (s2, 1), (s3, 1)])
    filename, data = excel_io.production_workbook(db, root.id)
    sheet = load_workbook(io.BytesIO(data)).active
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    assert root.name in filename
    assert len(rows) == 2
    assert sorted(r[5] for r in rows) == [3, 6]
    assert all(r[1] == a.code for r in rows)
    assert all('24V..蠕动泵 50%' in r[11] for r in rows)
    assert sheet.cell(2, 12).value == '选配组'
    assert 'L2:M2' in {str(r) for r in sheet.merged_cells.ranges}
    # Browser still returns backup rows; only Excel uses horizontal columns.
    assert any(r['is_backup_path'] for r in services.technical_bom(db, root.id, show_alternatives=True))
    _, hidden = excel_io.production_workbook(db, root.id, show_alternatives=False)
    hidden_sheet = load_workbook(io.BytesIO(hidden)).active
    assert hidden_sheet.max_column == 11
    assert sum(r[5] for r in hidden_sheet.iter_rows(min_row=4, values_only=True)) == 9
    # Selected candidate is first even if it was not first in member order.
    configure(db, s1, [(a, '50'), (b, '50')], selected=b)
    _, content = excel_io.technical_workbook(db, s1.id)
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet.cell(4, 3).value == b.code
    assert str(sheet.cell(4, 13).value).startswith(b.code)
    assert sheet.max_row == 4


def test_upper_alternatives_are_notes_not_counted_materials(db):
    a = make(db, '10.7781.0'); b = make(db, '10.7782.0')
    s1 = make(db, '05.781.01', 'semi_finished', [(a, 2)])
    s2 = make(db, '05.782.01', 'semi_finished', [(b, 7)])
    root = make(db, '00.781.01', 'machine', [(s1, 3)])
    configure(db, root, [(s1, '40'), (s2, '60')])
    _, data = excel_io.production_workbook(db, root.id)
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ['生产BOM', '上层选配说明']
    row = list(book.active.iter_rows(min_row=4, values_only=True))[0]
    assert row[1] == a.code and row[5] == 6
    assert book.active.max_column == 11  # No inherited parent alternatives mislabeled as leaf alternatives.
    assert book['上层选配说明'].cell(4, 5).value.startswith(s1.code)
    _, hidden = excel_io.production_workbook(db, root.id, show_alternatives=False)
    assert load_workbook(io.BytesIO(hidden)).sheetnames == ['生产BOM']


def test_excel_more_than_three_candidates_and_empty_cells(db):
    items = [make(db, f'10.79{i:02d}.0') for i in range(6)]
    root = make(db, '05.791.01', 'semi_finished', [(items[0], 1), (items[-1], 2)])
    configure(db, root, [(item, '20') for item in items[:5]])
    _, data = excel_io.technical_workbook(db, root.id)
    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet.max_column == 17
    assert all(sheet.cell(5, c).value is None for c in range(13, 18))
    assert sheet.cell(4, 7).value == 1 and sheet.cell(5, 7).value == 2


def test_source_transform_exclusion_zero_quantities_and_selected_only():
    records = {c: {'item_type': t} for c, t in [('J00.001.01', 'machine'), ('00.001.01', 'machine'),
        ('X03.001.01', 'semi_finished'), ('15.6015.0', 'material'), ('15.6014.0', 'material'), ('99.0001.9', 'material')]}
    def line(p, c, q):
        return {'parent_code': p, 'child_code': c, 'quantity': Decimal(q), 'row': 1, 'sheet': '测试'}
    rows = [line('J00.001.01', '99.0001.9', 1), line('00.001.01', 'X03.001.01', 3),
            line('X03.001.01', '15.6015.0', 0), line('X03.001.01', '15.6014.0', 0)]
    output, lines, report = transform_source(records, rows)
    assert 'J00.001.01' not in output and 'X03.001.01' not in output
    assert '99.0001.9' in output
    assert len(lines) == 1 and lines[0]['quantity'] == 3
    assert lines[0]['child_code'] == '15.6015.0'
    assert lines[0]['candidates'] == ['15.6015.0', '15.6014.0']
    assert len(report['zero_quantity_adjustments']) == 2
    assert rows[2]['quantity'] == 0  # Input is not mutated.
    with pytest.raises(ValueError):
        transform_source(records, [line('00.001.01', '15.6015.0', -1)])


def role_headers(db, role):
    user = models.User(username=role, password_hash=auth.hash_password('test-password'), display_name=role, role=role, active=True)
    db.add(user); db.commit()
    return {'Authorization': 'Bearer ' + auth.create_access_token(user)}


def test_imported_unofficial_codes_raise_allocation_floor_without_reserving_preview(db):
    pending = models.Item(item_type='material', code='99.0005.0', code_body='99.0005',
        version_label='0', version_number=0, version_series='99.0005', name='历史未正式',
        source_type='purchased', unit='pcs', status='active', is_formally_imported=False, unofficial_status='pending')
    db.add(pending); db.commit()
    assert services.preview_unofficial_code(db) == '99.0005.1'
    assert services.preview_unofficial_code(db) == '99.0005.1'
    assert services.allocate_unofficial_code(db) == '99.0005.1'
    assert services.allocate_unofficial_code(db) == '99.0005.2'
    db.commit()
    assert services.preview_unofficial_code(db) == '99.0005.3'


def test_guest_denied_all_exports_and_writes(client, db):
    headers = role_headers(db, 'guest')
    assert client.get('/api/items', headers=headers).status_code == 200
    for path in ['/api/items-basic-export?item_type=material', '/api/items/1/export/technical',
                 '/api/compare/export?item_ids=1', '/api/audit/export', '/api/imports/1/report']:
        assert client.get(path, headers=headers).status_code == 403, path
    for path, data in [('/api/items', {'item_type': 'material', 'name': '禁止', 'is_formally_imported': False}),
                       ('/api/bom/batch-export/technical', [1])]:
        assert client.post(path, headers=headers, json=data).status_code == 403
    assert client.get('/api/users', headers=headers).status_code == 403


def test_product_only_pending_archives_and_dev_user_management_guard(client, db, headers):
    product = role_headers(db, 'product'); dev = role_headers(db, 'dev')
    official = make(db, '10.7991.0')
    payload = {'item_type': 'material', 'name': '生产临时原材料', 'source_type': 'purchased',
               'is_formally_imported': False, 'similarity_confirmed': True}
    created = client.post('/api/items', headers=product, json=payload)
    assert created.status_code == 201, created.text
    item_id = created.json()['id']
    assert client.patch(f'/api/items/{item_id}', headers=product, json={'remark': '可编辑', 'reason': '测试'}).status_code == 200
    assert client.patch(f'/api/items/{official.id}', headers=product, json={'remark': '禁止', 'reason': '测试'}).status_code == 403
    assert client.delete(f'/api/items/{official.id}?reason=测试', headers=product).status_code == 403
    assert client.post('/api/items', headers=product, json={**payload, 'is_formally_imported': True, 'code': '10.7992.0'}).status_code == 403
    assert client.post('/api/items', headers=product, json={**payload, 'components': [{'child_item_id': official.id, 'quantity': 1}]}).status_code == 403
    assert client.post(f'/api/items/{item_id}/bom', headers=product, json={'child_item_id': official.id, 'quantity': 1, 'reason': '禁止'}).status_code == 403
    assert client.put(f'/api/items/{item_id}/path-alternatives', headers=product, json={}).status_code == 403
    assert client.post('/api/maintenance/apply', headers=product, json={'operation':'delete', 'parent_item_ids':[1], 'source_item_id':1, 'reason':'测试'}).status_code == 403
    assert client.get(f'/api/items/{official.id}/export/technical', headers=product).status_code == 200
    assert client.delete(f'/api/items/{item_id}?reason=测试', headers=product).status_code == 200
    assert client.get('/api/users', headers=dev).status_code == 403
    assert client.delete(f'/api/items/{official.id}?reason=测试', headers=dev).status_code == 200
    # Reset and rename never expose password hashes in API responses or audit snapshots.
    dev_id = db.scalar(select(models.User.id).where(models.User.role == 'dev'))
    response = client.patch(f'/api/users/{dev_id}', headers=headers, json={'username':'research', 'password':'reset-secret'})
    assert response.status_code == 200
    assert 'password' not in response.text
    assert client.post('/api/auth/login', json={'username':'research', 'password':'reset-secret'}).status_code == 200
    snapshots = '\n'.join(str(a.after_json) for a in db.scalars(select(models.AuditEvent)))
    assert 'reset-secret' not in snapshots and 'password_hash' not in snapshots
