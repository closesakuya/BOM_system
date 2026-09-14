from __future__ import annotations

import io
from urllib.parse import unquote

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DatabaseError


def create_item(client: TestClient, headers: dict[str, str], **overrides):
    payload = {
        "item_type": "material", "code": "10.9001.01", "name": "测试螺钉",
        "source_type": "purchased", "unit": "pcs", "reason": "自动化测试",
    }
    payload.update(overrides)
    response = client.post("/api/items", json=payload, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


def test_v1_1_nested_semi_finished_cycle_filter_replace_and_copy_delete(
    client: TestClient,
    headers: dict[str, str],
):
    material_a = create_item(client, headers, code="10.9701.01", name="嵌套测试原料甲")
    material_b = create_item(client, headers, code="10.9702.01", name="嵌套测试原料乙")
    material_c = create_item(client, headers, code="10.9703.01", name="嵌套测试原料丙")
    inner = create_item(
        client,
        headers,
        item_type="semi_finished",
        code="05.971.01",
        name="内部普通半成品",
        source_type=None,
        status="active",
        components=[{
            "child_item_id": material_a["id"], "quantity": "2.5", "sort_order": 7,
            "line_remark": "保留的行备注",
        }],
    )
    outer = create_item(
        client,
        headers,
        item_type="semi_finished",
        code="05.972.01",
        name="外部组合半成品",
        source_type=None,
        status="active",
        components=[{"child_item_id": inner["id"], "quantity": "3"}],
    )

    outer_detail = client.get(f"/api/items/{outer['id']}", headers=headers).json()
    assert outer_detail["is_combination"] is True
    combination = client.get(
        "/api/items",
        params={"item_type": "semi_finished", "semi_kind": "combination"},
        headers=headers,
    ).json()
    normal = client.get(
        "/api/items",
        params={"item_type": "semi_finished", "semi_kind": "normal"},
        headers=headers,
    ).json()
    assert {row["id"] for row in combination["items"]} == {outer["id"]}
    assert inner["id"] in {row["id"] for row in normal["items"]}

    cycle = client.post(
        f"/api/items/{inner['id']}/bom",
        headers=headers,
        json={"child_item_id": outer["id"], "quantity": "1", "reason": "构造环路应被阻止"},
    )
    assert cycle.status_code == 409
    assert "BOM 环路" in cycle.text
    assert "05.971.01" in cycle.text and "05.972.01" in cycle.text

    inner_detail = client.get(f"/api/items/{inner['id']}", headers=headers).json()
    original_line = inner_detail["components"][0]
    group = client.post(
        "/api/alternatives",
        headers=headers,
        json={
            "name": "替换清组验证",
            "item_type": "material",
            "default_item_id": material_a["id"],
            "member_item_ids": [material_a["id"], material_c["id"]],
            "reason": "创建替代组",
        },
    ).json()
    assigned = client.patch(
        f"/api/bom-lines/{original_line['id']}",
        headers=headers,
        json={"alternative_group_id": group["id"], "reason": "关联替代组"},
    )
    assert assigned.status_code == 200
    replaced = client.post(
        f"/api/bom-lines/{original_line['id']}/replace",
        headers=headers,
        json={"child_item_id": material_b["id"], "reason": "验证行级替换"},
    )
    assert replaced.status_code == 200, replaced.text
    replaced_row = replaced.json()
    assert replaced_row["child_item_id"] == material_b["id"]
    assert replaced_row["quantity"] == "2.500000"
    assert replaced_row["sort_order"] == 7
    assert replaced_row["line_remark"] == "保留的行备注"
    assert replaced_row["alternative_group_id"] is None
    assert replaced_row["alternative_group_cleared"] is True

    copied = client.post(
        f"/api/items/{inner['id']}/copy",
        headers=headers,
        json={"mode": "new_item", "code": "05.973.01", "reason": "复制删除边界测试"},
    )
    assert copied.status_code == 201, copied.text
    copied_item = copied.json()
    deleted_copy = client.delete(
        f"/api/items/{copied_item['id']}",
        params={"reason": "自身 BOM 和复制关系不应阻止删除"},
        headers=headers,
    )
    assert deleted_copy.status_code == 200, deleted_copy.text

    # Move the active source parent away from material B. The deleted copy
    # still retains its own outbound BOM and is now the only parent of B.
    moved_source = client.post(
        f"/api/bom-lines/{original_line['id']}/replace",
        headers=headers,
        json={"child_item_id": material_c["id"], "reason": "准备验证已删除父项引用边界"},
    )
    assert moved_source.status_code == 200, moved_source.text

    deleted_child = client.delete(
        f"/api/items/{material_b['id']}",
        params={"reason": "已删除父项不再构成有效引用"},
        headers=headers,
    )
    assert deleted_child.status_code == 200, deleted_child.text
    blocked_restore = client.post(
        f"/api/items/{copied_item['id']}/restore",
        params={"reason": "验证恢复检查"},
        headers=headers,
    )
    assert blocked_restore.status_code == 404
    assert client.get(f"/api/items/{material_b['id']}", headers=headers).status_code == 404


def test_v1_2_path_alternative_market_shares_and_bom_display_switches(client, headers):
    materials = [create_item(client, headers, code=f"10.980{i}.01", name=f"候选{i}") for i in range(1,3)]
    parent = create_item(client, headers, item_type="semi_finished", code="05.981.01",
                         name="选配父项", source_type=None, status="active",
                         components=[{"child_item_id":materials[0]["id"],"quantity":"2"}])
    line=client.get(f"/api/items/{parent['id']}/bom",headers=headers).json()[0]
    payload={"line_path":[line["id"]],"mode":"custom","selected_item_id":materials[0]["id"],
             "members":[{"item_id":materials[0]["id"],"market_share":"60.00"},
                        {"item_id":materials[1]["id"],"market_share":"40.00"}],"reason":"占比测试"}
    saved=client.put(f"/api/items/{parent['id']}/path-alternatives",headers=headers,json=payload)
    assert saved.status_code==200,saved.text
    expanded=client.get(f"/api/items/{parent['id']}/technical-bom",headers=headers).json()
    assert [r["market_share"] for r in expanded]==["60.00","40.00"]
    assert [r["is_backup_path"] for r in expanded]==[False,True]
    normal=client.get(f"/api/items/{parent['id']}/technical-bom",params={"show_alternatives":False},headers=headers).json()
    assert len(normal)==1
    production=client.get(f"/api/items/{parent['id']}/production-bom",headers=headers).json()
    assert [r["quantity"] for r in production]==["2","2"]
    payload["members"][1]["market_share"]="30.00"
    invalid=client.put(f"/api/items/{parent['id']}/path-alternatives",headers=headers,json=payload)
    assert invalid.status_code==422 and "100%" in invalid.text
    workbook=client.get(f"/api/items/{parent['id']}/export/technical",headers=headers)
    assert workbook.status_code==200


def test_login_dashboard_and_permissions(client: TestClient, headers: dict[str, str]):
    assert client.get("/api/dashboard", headers=headers).status_code == 200
    response = client.post("/api/users", headers=headers, json={
        "username": "viewer1", "password": "viewer1", "display_name": "查看者", "role": "viewer",
    })
    assert response.status_code == 201
    viewer_token = client.post("/api/auth/login", json={"username": "viewer1", "password": "viewer1"}).json()["access_token"]
    denied = client.post("/api/items", headers={"Authorization": f"Bearer {viewer_token}"}, json={
        "item_type": "material", "code": "10.9999.01", "name": "禁止写入", "source_type": "purchased",
    })
    assert denied.status_code == 403


def test_item_list_reports_total_and_pages_without_overlap(client: TestClient, headers: dict[str, str]):
    create_item(client, headers, code="10.9103.01", name="分页测试丙")
    create_item(client, headers, code="10.9101.01", name="分页测试甲")
    create_item(client, headers, code="10.9102.01", name="分页测试乙")

    first = client.get("/api/items", params={"item_type": "material", "limit": 2}, headers=headers).json()
    second = client.get(
        "/api/items", params={"item_type": "material", "limit": 2, "offset": 2}, headers=headers,
    ).json()

    assert first["total"] == second["total"] == 3
    assert [item["code"] for item in first["items"]] == ["10.9101.01", "10.9102.01"]
    assert [item["code"] for item in second["items"]] == ["10.9103.01"]
    assert {item["id"] for item in first["items"]}.isdisjoint(item["id"] for item in second["items"])


def test_item_list_filters_material_source_and_product_status(client: TestClient, headers: dict[str, str]):
    purchased = create_item(client, headers, code="10.9301.01", name="来源筛选外购物料")
    outsourced = create_item(
        client, headers, code="10.9302.01", name="来源筛选外协物料", source_type="outsourced",
    )
    self_made = create_item(
        client, headers, code="10.9303.01", name="来源筛选自制物料", source_type="self_made",
        components=[{"child_item_id": purchased["id"], "quantity": "1", "sort_order": 1}],
    )
    expected_sources = {
        "purchased": purchased["code"], "outsourced": outsourced["code"], "self_made": self_made["code"],
    }
    for source_type, code in expected_sources.items():
        result = client.get(
            "/api/items",
            params={"item_type": "material", "source_type": source_type, "q": "来源筛选", "limit": 1000},
            headers=headers,
        ).json()
        assert result["total"] == 1
        assert [item["code"] for item in result["items"]] == [code]
        assert all(item["source_type"] == source_type for item in result["items"])

    type_codes = {"semi_finished": "05.93", "unit": "03.93", "machine": "00.93"}
    statuses = {"trial": "试制", "active": "在用", "disabled": "停用"}
    for item_type, prefix in type_codes.items():
        expected: dict[str, str] = {}
        for index, (status, label) in enumerate(statuses.items(), start=1):
            item = create_item(
                client, headers, item_type=item_type, code=f"{prefix}{index}.01",
                name=f"{item_type}{label}筛选", source_type=None, status=status,
            )
            expected[status] = item["code"]
        for status, code in expected.items():
            result = client.get(
                "/api/items",
                params={"item_type": item_type, "status": status, "limit": 1000},
                headers=headers,
            ).json()
            assert result["total"] == 1
            assert [item["code"] for item in result["items"]] == [code]
            assert all(item["status"] == status for item in result["items"])


def test_v1_1_unofficial_material_model_references_and_basic_export(
    client: TestClient,
    headers: dict[str, str],
):
    official = create_item(
        client,
        headers,
        code="10.9501.01",
        name="正式分析面板",
        key_component_code="B01.A",
    )
    unofficial_a = create_item(
        client,
        headers,
        code="10.9501.01",
        name="临时分析面板甲",
        is_formally_imported=False,
        similarity_confirmed=True,
        components=[{
            "child_item_id": official["id"], "quantity": "2.5", "sort_order": 3,
            "line_remark": "待转正式组成",
        }],
    )
    unofficial_b = create_item(
        client,
        headers,
        code="10.9501.01",
        name="临时分析面板乙",
        is_formally_imported=False,
        similarity_confirmed=True,
    )
    assert unofficial_a["code"] == "99.0001.0"
    assert unofficial_b["code"] == "99.0001.1"
    assert unofficial_a["unofficial_status"] == unofficial_b["unofficial_status"] == "pending"

    filtered = client.get(
        "/api/items",
        params={"item_type": "material", "is_formally_imported": False, "q": "99.0001"},
        headers=headers,
    ).json()
    assert filtered["total"] == 2
    assert {item["id"] for item in filtered["items"]} == {unofficial_a["id"], unofficial_b["id"]}
    availability = client.get(
        "/api/items/code-availability",
        params={
            "code": "ignored-by-server",
            "item_type": "material",
            "is_formally_imported": False,
        },
        headers=headers,
    ).json()
    assert availability["exists"] is False
    assert availability["code"] == "99.0001.2"
    assert availability["duplicate_items"] == []

    semi = create_item(
        client,
        headers,
        item_type="semi_finished",
        code="05.951.01",
        name="机型引用半成品",
        source_type=None,
        status="active",
        components=[{"child_item_id": official["id"], "quantity": "1"}],
    )
    machine = create_item(
        client,
        headers,
        item_type="machine",
        code="00.951.01",
        name="R60测试整机",
        source_type=None,
        status="active",
        machine_model="R60",
        components=[{"child_item_id": semi["id"], "quantity": "1"}],
    )
    assert machine["machine_model"] == "R60"
    ownership = client.get(f"/api/items/{official['id']}", headers=headers).json()["ownership"]
    assert ownership["current_models"][0]["model"] == "R60"
    assert ownership["current_models"][0]["machines"][0]["id"] == machine["id"]
    assert client.get("/api/machine-models", headers=headers).json() == ["R60"]

    denied_bom = client.post(
        f"/api/items/{semi['id']}/bom",
        headers=headers,
        json={"child_item_id": unofficial_a["id"], "quantity": "1", "reason": "禁止临时物料引用"},
    )
    assert denied_bom.status_code == 409
    denied_group = client.post(
        "/api/alternatives",
        headers=headers,
        json={
            "name": "临时物料禁止替代",
            "item_type": "material",
            "default_item_id": official["id"],
            "member_item_ids": [unofficial_a["id"]],
            "reason": "验证未正式限制",
        },
    )
    assert denied_group.status_code == 409

    rule = client.post(
        "/api/code-rules",
        headers=headers,
        json={
            "item_type": "material", "large_category": "测试大类", "small_category": "测试小类",
            "prefix": "10.", "pattern": "10.XXXX.XX", "active": True,
        },
    ).json()
    promotion_payload = {
        "code_rule_id": rule["id"], "auxiliary_code": unofficial_a.get("auxiliary_code"),
        "name": unofficial_a["name"], "specification": unofficial_a.get("specification"),
        "source_type": unofficial_a["source_type"], "unit": unofficial_a["unit"],
        "remark": unofficial_a.get("remark"), "previous_version_name": unofficial_a.get("previous_version_name"),
        "invoice_name": unofficial_a.get("invoice_name"), "material_attribute": unofficial_a.get("material_attribute"),
        "key_component_code": "b01.a", "similarity_confirmed": True,
    }
    conflict = client.post(
        f"/api/items/{unofficial_a['id']}/promote",
        headers=headers,
        json={**promotion_payload, "code": "10.9501.01", "reason": "冲突转正式"},
    )
    assert conflict.status_code == 409
    promoted = client.post(
        f"/api/items/{unofficial_a['id']}/promote",
        headers=headers,
        json={**promotion_payload, "code": "10.9502.01", "reason": "确认转正式"},
    )
    assert promoted.status_code == 200, promoted.text
    assert promoted.json()["code"] == "10.9502.01"
    assert promoted.json()["is_formally_imported"] is True
    assert promoted.json()["id"] != unofficial_a["id"]
    assert promoted.json()["key_component_code"] == "B01.A"
    archived = client.get(f"/api/items/{unofficial_a['id']}", headers=headers).json()
    assert archived["code"] == "99.0001.0"
    assert archived["unofficial_status"] == "archived"
    assert archived["promotion_trace"]["promoted_to"]["id"] == promoted.json()["id"]
    target = client.get(f"/api/items/{promoted.json()['id']}", headers=headers).json()
    assert [row["id"] for row in target["promotion_trace"]["promoted_from"]] == [unofficial_a["id"]]
    assert len(target["components"]) == 1
    assert target["components"][0]["child_item_id"] == official["id"]
    assert target["components"][0]["quantity"] == "2.500000"
    assert target["components"][0]["sort_order"] == 3
    assert target["components"][0]["line_remark"] == "待转正式组成"
    references_after = client.get(f"/api/items/{official['id']}/references", headers=headers).json()
    assert promoted.json()["id"] in {row["parent"]["id"] for row in references_after}
    assert unofficial_a["id"] not in {row["parent"]["id"] for row in references_after}
    blocked_edit = client.patch(
        f"/api/items/{unofficial_a['id']}", headers=headers,
        json={"name": "不得修改", "reason": "验证封存保护"},
    )
    assert blocked_edit.status_code == 409
    blocked_delete = client.delete(
        f"/api/items/{unofficial_a['id']}", headers=headers, params={"reason": "验证封存保护"},
    )
    assert blocked_delete.status_code == 409

    exported = client.get(
        "/api/items-basic-export",
        params={"item_type": "material", "q": "分析面板"},
        headers=headers,
    )
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [cell.value for cell in sheet[1]] == [
        "序号","关键器件码","物料编码","*名称","*规格型号","*属性","*单位","备注",
        "上一版本名称","发票名称","*材料属性","变更记录","当前在用归属产品机型"]
    exported_codes = {row[2] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert {"10.9501.01", "99.0001.0", "99.0001.1", "10.9502.01"} <= exported_codes


def test_item_search_queries_all_rows_and_prioritizes_name_contains(
    client: TestClient, headers: dict[str, str],
):
    create_item(client, headers, code="17.2001.01", name="搜索测试乙")
    create_item(client, headers, code="17.1001.01", name="搜索测试甲")
    create_item(client, headers, code="10.9201.01", name="名称含17.的物料", specification="特殊规格Z17")

    prefix = client.get("/api/items", params={"q": "17.", "limit": 2}, headers=headers).json()
    assert prefix["total"] == 3
    assert [item["code"] for item in prefix["items"]] == ["10.9201.01", "17.1001.01"]

    exact = client.get("/api/items", params={"q": "17.2001.01"}, headers=headers).json()
    assert exact["items"][0]["code"] == "17.2001.01"
    specification = client.get("/api/items", params={"q": "特殊规格Z17"}, headers=headers).json()
    assert specification["total"] == 1
    assert specification["items"][0]["code"] == "10.9201.01"


def test_live_code_availability_and_two_recommendations(client: TestClient, headers: dict[str, str]):
    create_item(client, headers, code="10.1001.01", name="编码校验源物料")
    create_item(client, headers, code="10.1012.01", name="编码校验系列最大物料")

    conflict = client.get("/api/items/code-availability", params={"code": "10.1001.01"}, headers=headers)
    assert conflict.status_code == 200
    result = conflict.json()
    assert result["valid"] is True
    assert result["exists"] is True
    assert result["maximum"]["code"] == "10.1012.01"
    assert result["next_body_code"] == "10.1013.01"
    assert result["next_version_code"] == "10.1001.02"

    available = client.get("/api/items/code-availability", params={"code": "10.1001.02"}, headers=headers).json()
    assert available["valid"] is True
    assert available["exists"] is False
    assert available["next_body_code"] == "10.1013.01"
    assert available["next_version_code"] == "10.1001.02"

    invalid = client.get("/api/items/code-availability", params={"code": "101001"}, headers=headers).json()
    assert invalid["valid"] is False
    assert "版本尾号" in invalid["message"]


def test_selected_code_rule_controls_availability_and_creation(client: TestClient, headers: dict[str, str], db):
    from backend.app import models

    rule = models.CodeRule(
        item_type="material", large_category="螺钉紧固,管材标准件", small_category="螺钉",
        material_attribute="螺钉", prefix="19.1",
        pattern="19.1XXX.0  例如：19.1001.0", active=True, sort_order=1,
    )
    db.add(rule)
    db.commit()

    invalid = client.get(
        "/api/items/code-availability", params={"code": "10.01", "code_rule_id": rule.id}, headers=headers,
    ).json()
    assert invalid["valid"] is False
    assert "19.1XXX.0" in invalid["message"]

    valid = client.get(
        "/api/items/code-availability", params={"code": "19.1001.0", "code_rule_id": rule.id}, headers=headers,
    ).json()
    assert valid["valid"] is True

    rejected = client.post("/api/items", headers=headers, json={
        "item_type": "material", "code": "10.01", "code_rule_id": rule.id,
        "name": "错误规则编码", "source_type": "purchased", "unit": "pcs", "reason": "规则校验测试",
    })
    assert rejected.status_code == 422

    recommended = client.get(f"/api/code-rules/{rule.id}/recommend", headers=headers).json()
    assert recommended["recommended_code"] == "19.1001.0"


def test_default_non_material_rules_control_all_three_item_types(
    client: TestClient, headers: dict[str, str], db,
):
    from backend.app import services

    assert services.ensure_default_non_material_code_rules(db) == 3
    assert services.ensure_default_non_material_code_rules(db) == 0
    rules = client.get("/api/code-rules", headers=headers).json()
    expected = {
        "semi_finished": ("半成品", "模块和组件（CBB）", "05.", "05.XXX.01"),
        "unit": ("单元", "分析单元或电气机柜", "03.", "03.XXX.01"),
        "machine": ("整机", "包装后的成品机", "00.", "00.XXX.01"),
    }
    for item_type, values in expected.items():
        rule = next(row for row in rules if row["item_type"] == item_type)
        assert (rule["large_category"], rule["small_category"], rule["prefix"], rule["pattern"]) == values
        valid_code = {"semi_finished": "05.901.01", "unit": "03.901.01", "machine": "00.901.01"}[item_type]
        valid = client.get(
            "/api/items/code-availability",
            params={"code": valid_code, "code_rule_id": rule["id"], "item_type": item_type},
            headers=headers,
        ).json()
        assert valid["valid"] is True
        invalid = client.get(
            "/api/items/code-availability",
            params={"code": "10.1001.01", "code_rule_id": rule["id"], "item_type": item_type},
            headers=headers,
        ).json()
        assert invalid["valid"] is False

    unit_rule = next(row for row in rules if row["item_type"] == "unit")
    rejected = client.post("/api/items", headers=headers, json={
        "item_type": "machine", "code": "00.902.01", "code_rule_id": unit_rule["id"],
        "name": "跨类型规则整机", "source_type": None, "unit": "pcs", "reason": "规则类型校验",
    })
    assert rejected.status_code == 422


def test_item_bom_expansion_cycle_copy_and_exports(client: TestClient, headers: dict[str, str]):
    material = create_item(client, headers)
    nested = create_item(client, headers, code="10.9002.01", name="外协组件", source_type="outsourced", components=[{
        "child_item_id": material["id"], "quantity": "2.5", "sort_order": 1,
    }])
    semi = create_item(client, headers, item_type="semi_finished", code="05.901.01", name="测试半成品", source_type=None, components=[{
        "child_item_id": nested["id"], "quantity": "2", "sort_order": 1,
    }])
    unit = create_item(client, headers, item_type="unit", code="03.901.01", name="测试单元", source_type=None, components=[{
        "child_item_id": semi["id"], "quantity": "3", "sort_order": 1,
    }])
    machine = create_item(client, headers, item_type="machine", code="00.901.01", name="测试整机", source_type=None, components=[{
        "child_item_id": unit["id"], "quantity": "4", "sort_order": 1,
    }])
    production = client.get(f"/api/items/{machine['id']}/production-bom", headers=headers)
    assert production.status_code == 200
    assert production.json()[0]["item"]["code"] == nested["code"]
    assert production.json()[0]["quantity"] == "24"
    expanded_production = client.get(
        f"/api/items/{machine['id']}/production-bom",
        params={"expand_materials": True}, headers=headers,
    )
    assert expanded_production.json()[0]["item"]["code"] == material["code"]
    assert expanded_production.json()[0]["quantity"] == "60"
    technical = client.get(f"/api/items/{machine['id']}/technical-bom", headers=headers)
    assert len(technical.json()) == 3
    references = client.get(f"/api/items/{material['id']}/references", headers=headers)
    assert references.status_code == 200
    direct = next(row for row in references.json() if row["is_direct"])
    assert [item["code"] for item in direct["path_items"]] == [nested["code"], material["code"]]
    assert all(isinstance(item["name"], str) for item in direct["path_items"])

    parent_candidates = client.get(
        f"/api/maintenance/parent-candidates?operation=delete&source_item_id={material['id']}", headers=headers,
    )
    assert parent_candidates.status_code == 200
    assert [item["id"] for item in parent_candidates.json()["items"]] == [nested["id"]]
    add_candidates = client.get(
        f"/api/maintenance/parent-candidates?operation=add&target_item_id={semi['id']}", headers=headers,
    )
    assert add_candidates.status_code == 200
    assert unit["id"] in [item["id"] for item in add_candidates.json()["items"]]
    assert {item["item_type"] for item in add_candidates.json()["items"]} <= {"unit", "machine"}

    unit_line = client.get(f"/api/items/{unit['id']}/bom", headers=headers).json()[0]
    quantity_update = client.patch(f"/api/bom-lines/{unit_line['id']}", headers=headers, json={
        "quantity": "4.5", "reason": "验证 BOM 行数量修改",
    })
    assert quantity_update.status_code == 200, quantity_update.text
    assert client.get(f"/api/items/{unit['id']}", headers=headers).json()["components"][0]["quantity"] == "4.500000"
    assert any(event["reason"] == "验证 BOM 行数量修改" for event in client.get(f"/api/items/{unit['id']}/history", headers=headers).json())

    material_update = client.patch(f"/api/items/{material['id']}", headers=headers, json={
        "name": "测试螺钉（新名称）", "specification": "M4-新规格", "remark": "新备注",
        "reason": "验证子物料档案同步", "similarity_confirmed": True,
    })
    assert material_update.status_code == 200, material_update.text
    nested_detail = client.get(f"/api/items/{nested['id']}", headers=headers).json()
    assert nested_detail["components"][0]["child"]["name"] == "测试螺钉（新名称）"
    refreshed_technical = client.get(
        f"/api/items/{machine['id']}/technical-bom",
        params={"expand_materials": True}, headers=headers,
    ).json()
    refreshed_material = next(row["item"] for row in refreshed_technical if row["item"]["id"] == material["id"])
    assert (refreshed_material["name"], refreshed_material["specification"], refreshed_material["remark"]) == ("测试螺钉（新名称）", "M4-新规格", "新备注")
    material_history = client.get(f"/api/items/{material['id']}/history", headers=headers).json()
    sync_event = next(event for event in material_history if event["reason"] == "验证子物料档案同步")
    assert {change["field"] for change in sync_event["changes"]} >= {"name", "specification", "remark"}
    ancestor_history = client.get(f"/api/items/{machine['id']}/history", headers=headers).json()
    descendant_event = next(event for event in ancestor_history if event["reason"] == "验证子物料档案同步")
    assert descendant_event["scope"] == "descendant"
    assert descendant_event["subject"]["code"] == material["code"]
    assert [item["code"] for item in descendant_event["paths"][0]] == [
        machine["code"], unit["code"], semi["code"], nested["code"], material["code"],
    ]

    removable = create_item(client, headers, code="10.9003.01", name="待移除组件")
    added_line = client.post(f"/api/items/{semi['id']}/bom", headers=headers, json={
        "child_item_id": removable["id"], "quantity": "1", "reason": "验证新增组件名称编码",
    })
    assert added_line.status_code == 201, added_line.text
    removed_line_id = added_line.json()["id"]
    removed = client.delete(
        f"/api/bom-lines/{removed_line_id}?reason=验证移除组件名称编码", headers=headers,
    )
    assert removed.status_code == 200, removed.text
    relation_history = client.get(f"/api/items/{machine['id']}/history", headers=headers).json()
    add_event = next(event for event in relation_history if event["reason"] == "验证新增组件名称编码")
    remove_event = next(event for event in relation_history if event["reason"] == "验证移除组件名称编码")
    assert add_event["subject"]["code"] == removable["code"]
    assert add_event["changes"][0] == {"field": "component", "before": None, "after": "10.9003.01｜待移除组件"}
    assert remove_event["subject"]["code"] == removable["code"]
    assert remove_event["changes"][0] == {"field": "component", "before": "10.9003.01｜待移除组件", "after": None}
    assert [item["code"] for item in remove_event["paths"][0]][-2:] == [semi["code"], removable["code"]]

    updated = client.patch(f"/api/items/{machine['id']}", headers=headers, json={
        "name": "测试整机（已变更）",
        "reason": "验证物料变更记录", "similarity_confirmed": True,
    })
    assert updated.status_code == 200, updated.text
    history = client.get(f"/api/items/{machine['id']}/history", headers=headers)
    assert history.status_code == 200, history.text
    update_event = next(event for event in history.json() if event["reason"] == "验证物料变更记录")
    assert any(change["field"] == "name" for change in update_event["changes"])

    cycle = client.post(f"/api/items/{nested['id']}/bom", headers=headers, json={
        "child_item_id": nested["id"], "quantity": 1, "reason": "环路验证",
    })
    assert cycle.status_code == 409
    copied = client.post(f"/api/items/{machine['id']}/copy", headers=headers, json={"mode": "new_version", "reason": "版本测试"})
    assert copied.status_code == 201, copied.text
    assert copied.json()["code"].endswith(".02")

    direct_material = create_item(client, headers, code="10.9004.01", name="整机直接材料")
    direct_line = client.post(f"/api/items/{machine['id']}/bom", headers=headers, json={
        "child_item_id": direct_material["id"], "quantity": "1", "sort_order": 2,
        "reason": "验证技术 BOM 树形层级",
    })
    assert direct_line.status_code == 201, direct_line.text

    for kind in ("technical", "production"):
        exported = client.get(
            f"/api/items/{machine['id']}/export/{kind}",
            params={"expand_materials": True}, headers=headers,
        )
        assert exported.status_code == 200
        assert exported.content[:2] == b"PK"
        exported_name = unquote(exported.headers["content-disposition"].split("UTF-8''", 1)[1])
        expected_suffix = "技术BOM.xlsx" if kind == "technical" else "生产BOM.xlsx"
        assert exported_name == f"{machine['code']}-测试整机（已变更）-{expected_suffix}"
        sheet = load_workbook(io.BytesIO(exported.content)).active
        code_column = 3 if kind == "technical" else 2
        material_row = next(row for row in sheet.iter_rows(min_row=3, values_only=True) if row[code_column - 1] == material["code"])
        name_column = 4 if kind == "technical" else 3
        specification_column = 5 if kind == "technical" else 4
        remark_column = 8 if kind == "technical" else 7
        assert material_row[name_column - 1] == "测试螺钉（新名称）"
        assert material_row[specification_column - 1] == "M4-新规格"
        assert material_row[remark_column - 1] == "新备注"
        if kind == "technical":
            rows = list(sheet.iter_rows(min_row=3, values_only=True))
            assert [row[1] for row in rows] == [
                "├─ 1级", "│  └─ 2级", "│     └─ 3级", "│        └─ 4级", "└─ 1级",
            ]
            assert sheet.column_dimensions["B"].width == 32
            assert all(cell.alignment.wrap_text is not True for cell in sheet["B"][2:])


def test_similarity_delete_guard_and_bulk_maintenance(client: TestClient, headers: dict[str, str]):
    first = create_item(client, headers)
    conflict = client.post("/api/items", headers=headers, json={
        "item_type": "material", "code": "10.9002.01", "name": "测试螺钉",
        "source_type": "purchased", "unit": "pcs", "reason": "相似验证",
    })
    assert conflict.status_code == 409
    second = create_item(client, headers, code="10.9002.01", name="测试螺钉", similarity_confirmed=True)
    third = create_item(client, headers, code="10.9003.01", name="测试螺母")
    parent = create_item(client, headers, item_type="semi_finished", code="05.902.01", name="批量维护对象", source_type=None)
    applied = client.post("/api/maintenance/apply", headers=headers, json={
        "operation": "add", "target_item_id": first["id"], "parent_item_ids": [parent["id"]],
        "quantity": "2", "reason": "批量新增",
    })
    assert applied.status_code == 200, applied.text
    preview_merge = client.post("/api/maintenance/preview", headers=headers, json={
        "operation": "add", "target_item_id": first["id"], "parent_item_ids": [parent["id"]], "quantity": "3",
    })
    assert preview_merge.status_code == 200, preview_merge.text
    assert preview_merge.json()["lines"][0]["effect"] == "merge"
    assert preview_merge.json()["lines"][0]["before_quantity"] == "2.000000"
    assert preview_merge.json()["lines"][0]["after_quantity"] == "5.000000"
    replaced = client.post("/api/maintenance/apply", headers=headers, json={
        "operation": "replace", "source_item_id": first["id"], "target_item_id": second["id"],
        "parent_item_ids": [parent["id"]], "reason": "批量替换",
    })
    assert replaced.status_code == 200, replaced.text
    group = client.post("/api/alternatives", headers=headers, json={
        "name": "测试替代组", "item_type": "material", "default_item_id": second["id"],
        "member_item_ids": [first["id"]], "reason": "建立替代组",
    })
    assert group.status_code == 201, group.text
    line = client.get(f"/api/items/{parent['id']}/bom", headers=headers).json()[0]
    assigned = client.patch(f"/api/bom-lines/{line['id']}", headers=headers, json={
        "alternative_group_id": group.json()["id"], "reason": "关联替代组",
    })
    assert assigned.status_code == 200, assigned.text
    switched = client.post(f"/api/bom-lines/{line['id']}/alternative-selection", headers=headers, json={
        "item_id": first["id"], "reason": "切换当前替代件",
    })
    assert switched.status_code == 200, switched.text
    assert switched.json()["child_item_id"] == first["id"]
    assert switched.json()["child"]["code"] == first["code"]
    assert switched.json()["quantity"] == "2.000000"
    blocked_member_removal = client.patch(f"/api/alternatives/{group.json()['id']}", headers=headers, json={
        "name": "测试替代组", "item_type": "material", "default_item_id": second["id"],
        "member_item_ids": [third["id"]], "active": True, "reason": "错误移除当前选中成员",
    })
    assert blocked_member_removal.status_code == 409
    assert first["code"] in blocked_member_removal.text
    edited_group = client.patch(f"/api/alternatives/{group.json()['id']}", headers=headers, json={
        "name": "测试替代组-已修改", "item_type": "material", "default_item_id": first["id"],
        "member_item_ids": [second["id"], third["id"]], "active": True, "reason": "验证替代组编辑",
    })
    assert edited_group.status_code == 200, edited_group.text
    assert edited_group.json()["name"] == "测试替代组-已修改"
    assert edited_group.json()["default_item_id"] == first["id"]
    assert {member["item_id"] for member in edited_group.json()["members"]} == {first["id"], second["id"], third["id"]}
    assert client.get(f"/api/items/{parent['id']}/bom", headers=headers).json()[0]["child_item_id"] == first["id"]

    # A target already present in the same parent's direct BOM is rejected.
    # Neither the selected line nor the existing target quantity may change.
    direct_target = client.post(f"/api/items/{parent['id']}/bom", headers=headers, json={
        "child_item_id": second["id"], "quantity": "5", "reason": "建立已存在目标行",
    })
    assert direct_target.status_code == 201, direct_target.text
    switched_to_existing = client.post(f"/api/bom-lines/{line['id']}/alternative-selection", headers=headers, json={
        "item_id": second["id"], "reason": "切换到已存在的替代件",
    })
    assert switched_to_existing.status_code == 409, switched_to_existing.text
    assert parent["code"] in switched_to_existing.text
    assert second["code"] in switched_to_existing.text
    assert "不会合并数量" in switched_to_existing.text
    final_lines = client.get(f"/api/items/{parent['id']}/bom", headers=headers).json()
    assert next(row for row in final_lines if row["id"] == line["id"])["child_item_id"] == first["id"]
    assert next(row for row in final_lines if row["id"] == line["id"])["quantity"] == "2.000000"
    assert next(row for row in final_lines if row["child_item_id"] == second["id"])["quantity"] == "5.000000"

    # The same material deeper in a semi-finished subtree is not a direct
    # sibling conflict. A machine-level alternative may still select it.
    nested_semi = create_item(
        client, headers, item_type="semi_finished", code="05.903.01",
        name="包含目标物料的下级半成品", source_type=None,
        components=[{"child_item_id": second["id"], "quantity": "3"}],
    )
    machine = create_item(
        client, headers, item_type="machine", code="00.903.01",
        name="跨层选配验证整机", source_type=None,
        components=[
            {"child_item_id": nested_semi["id"], "quantity": "1", "sort_order": 1},
            {"child_item_id": third["id"], "quantity": "4", "sort_order": 2},
        ],
    )
    machine_lines = client.get(f"/api/items/{machine['id']}/bom", headers=headers).json()
    machine_option_line = next(row for row in machine_lines if row["child_item_id"] == third["id"])
    assert machine_option_line["alternative_group_id"] is None
    indirect_allowed = client.post(
        f"/api/bom-lines/{machine_option_line['id']}/alternative-selection", headers=headers,
        json={
            "item_id": second["id"], "alternative_group_id": group.json()["id"],
            "reason": "自动匹配替代组；下级已有目标但直接层允许选配",
        },
    )
    assert indirect_allowed.status_code == 200, indirect_allowed.text
    assert indirect_allowed.json()["child_item_id"] == second["id"]
    assert indirect_allowed.json()["quantity"] == "4.000000"
    assert indirect_allowed.json()["alternative_group_id"] == group.json()["id"]
    assert len([row for row in client.get(f"/api/items/{machine['id']}/bom", headers=headers).json() if row["child_item_id"] == second["id"]]) == 1
    alternative_audit = client.get(
        f"/api/audit?entity_type=alternative_group&entity_id={group.json()['id']}", headers=headers,
    ).json()
    assert any(event["reason"] == "验证替代组编辑" and event["action"] == "update" for event in alternative_audit)
    denied = client.delete(f"/api/items/{first['id']}?reason=引用保护", headers=headers)
    assert denied.status_code == 409


def test_material_assembly_rules_and_leaf_compatibility(client: TestClient, headers: dict[str, str]):
    self_made_without_components = client.post("/api/items", headers=headers, json={
        "item_type": "material", "code": "10.9005.01", "name": "无来源组件",
        "source_type": "self_made", "unit": "pcs", "reason": "规则验证",
    })
    assert self_made_without_components.status_code == 201
    assert self_made_without_components.json()["requires_assembly"] is False

    leaf = create_item(
        client, headers, code="11.9006.01", name="外协叶子原材料", source_type="outsourced",
    )
    assert leaf["requires_assembly"] is False
    assert client.get(f"/api/items/{leaf['id']}/technical-bom", headers=headers).json() == []
    leaf_export = client.get(f"/api/items/{leaf['id']}/export/technical", headers=headers)
    assert leaf_export.status_code == 200
    leaf_sheet = load_workbook(io.BytesIO(leaf_export.content)).active
    assert leaf_sheet.cell(3, 3).value == leaf["code"]
    updated_leaf = client.patch(f"/api/items/{leaf['id']}", headers=headers, json={
        "specification": "图号-LEAF-02", "remark": "无组成，按普通原材料维护",
        "reason": "验证外协叶子档案修改", "similarity_confirmed": True,
    })
    assert updated_leaf.status_code == 200, updated_leaf.text
    assert updated_leaf.json()["specification"] == "图号-LEAF-02"
    assert updated_leaf.json()["remark"] == "无组成，按普通原材料维护"
    assert updated_leaf.json()["requires_assembly"] is False

    purchased_leaf = create_item(client, headers, code="10.9007.01", name="外购基础件")
    semi = create_item(
        client, headers, item_type="semi_finished", code="05.907.01", name="叶子展开验证",
        source_type=None, components=[{"child_item_id": leaf["id"], "quantity": "2"}],
    )
    production_before = client.get(f"/api/items/{semi['id']}/production-bom", headers=headers).json()
    assert production_before[0]["item"]["code"] == leaf["code"]

    candidates = client.get(
        f"/api/maintenance/parent-candidates?operation=add&target_item_id={purchased_leaf['id']}",
        headers=headers,
    ).json()["items"]
    assert leaf["id"] in [item["id"] for item in candidates]

    added = client.post(f"/api/items/{leaf['id']}/bom", headers=headers, json={
        "child_item_id": purchased_leaf["id"], "quantity": "3", "reason": "外协叶子增加组成",
    })
    assert added.status_code == 201, added.text
    assert client.get(f"/api/items/{leaf['id']}", headers=headers).json()["requires_assembly"] is True
    blocked_conversion = client.patch(f"/api/items/{leaf['id']}", headers=headers, json={
        "source_type": "purchased", "reason": "有组成时禁止改外购", "similarity_confirmed": True,
    })
    assert blocked_conversion.status_code == 409
    production_after = client.get(
        f"/api/items/{semi['id']}/production-bom",
        params={"expand_materials": True}, headers=headers,
    ).json()
    assert production_after[0]["item"]["code"] == purchased_leaf["code"]
    assert production_after[0]["quantity"] == "6"

    removed = client.delete(
        f"/api/bom-lines/{added.json()['id']}?reason=移除最后一个组成", headers=headers,
    )
    assert removed.status_code == 200, removed.text
    leaf_after_removal = client.get(f"/api/items/{leaf['id']}", headers=headers).json()
    assert leaf_after_removal["requires_assembly"] is False
    converted = client.patch(f"/api/items/{leaf['id']}", headers=headers, json={
        "source_type": "purchased", "reason": "无组成时改为外购", "similarity_confirmed": True,
    })
    assert converted.status_code == 200, converted.text

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["物料编码", "名称", "规格型号", "属性", "单位", "备注"])
    sheet.append(["11.9008.01", "Excel外协叶子", "OUT-LEAF", "外协", "pcs", "无组成导入"])
    content = io.BytesIO(); workbook.save(content)
    preview = client.post(
        "/api/imports/materials/preview", headers=headers,
        files={"file": ("outsourced-leaf.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 200, preview.text
    imported_row = preview.json()["rows"][0]
    assert imported_row["status"] in {"ready", "yellow"}
    assert imported_row["error_message"] is None
    committed = client.post(
        f"/api/imports/{imported_row['batch_id']}/commit", headers=headers,
        json={"row_ids": [imported_row["id"]], "similarity_confirmed_row_ids": [], "reason": "导入外协叶子"},
    )
    assert committed.status_code == 200, committed.text
    assert committed.json()["created"][0]["source_type"] == "outsourced"
    assert committed.json()["created"][0]["requires_assembly"] is False


def test_material_import_rows_are_editable_and_commit_atomically(
    client: TestClient,
    headers: dict[str, str],
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["系统物料编码", "名称", "规格型号", "属性", "单位", "关键器件码", "导入状态"])
    sheet.append(["11.9601.01", "正式导入测试件", "FORMAL-01", "外购", "pcs", "B02.A", "正式导入"])
    sheet.append(["LEGACY-TEMP-99", "临时导入测试件", "TEMP-99", "自制", "pcs", "B02.B", ""])
    content = io.BytesIO(); workbook.save(content)
    preview = client.post(
        "/api/imports/materials/preview",
        headers=headers,
        files={"file": ("v11-import.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 200, preview.text
    batch = preview.json()["batch"]
    rows = preview.json()["rows"]
    assert len(rows) == 2
    assert any(row["status"] == "error" for row in rows)

    partial = client.post(
        f"/api/imports/{batch['id']}/commit",
        headers=headers,
        json={"row_ids": [rows[0]["id"]], "similarity_confirmed_row_ids": [], "reason": "禁止部分导入"},
    )
    assert partial.status_code == 409

    second = rows[1]
    edited_payload = {
        **second["payload"],
        "is_formally_imported": False,
        "cancelled": False,
    }
    edited = client.patch(
        f"/api/imports/{batch['id']}/rows/{second['id']}",
        headers=headers,
        json=edited_payload,
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["status"] in {"ready", "yellow", "red"}

    refreshed = client.get(f"/api/imports/{batch['id']}", headers=headers).json()
    active = [row for row in refreshed["rows"] if row["status"] != "cancelled"]
    confirmed = [row["id"] for row in active if row["status"] == "red"]
    committed = client.post(
        f"/api/imports/{batch['id']}/commit",
        headers=headers,
        json={
            "row_ids": [row["id"] for row in active],
            "similarity_confirmed_row_ids": confirmed,
            "reason": "整批原子导入",
        },
    )
    assert committed.status_code == 200, committed.text
    assert committed.json()["batch_status"] == "completed"
    created = committed.json()["created"]
    assert {row["code"] for row in created} == {"11.9601.01", "99.0001.0"}
    temporary = next(row for row in created if not row["is_formally_imported"])
    assert temporary["source_type"] == "self_made"
    assert temporary["historical_item_code"] == "LEGACY-TEMP-99"
    assert temporary["key_component_code"] == "B02.B"


def test_code_rule_controlled_replace_snapshot_restore_and_immutable_audit(client: TestClient, headers: dict[str, str], db):
    from backend.app import services

    services.ensure_default_non_material_code_rules(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表2 原材料分类及编码规则表"
    sheet.append(["原材料分类及物料编码规则表"])
    sheet.append(["序号", "大类", "小类", "材料属性", "编码开头", "编码示例(XXX/XX表示ID序列号）", "说明"])
    sheet.append([1, "测试大类", "测试小类", "测试属性", "98.7", "98.7XXX.01", "自动化测试"])
    content = io.BytesIO(); workbook.save(content)
    preview = client.post("/api/code-rule-imports/preview", headers=headers, files={"file": ("rules.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 200, preview.text
    assert preview.json()["counts"]["rules"] == 1
    assert preview.json()["counts"]["removed"] == 0
    applied = client.post("/api/code-rule-imports/apply", headers=headers, json={"rules": preview.json()["rules"], "reason": "规则覆盖测试"})
    assert applied.status_code == 200, applied.text
    remaining = client.get("/api/code-rules", headers=headers).json()
    assert len([row for row in remaining if row["item_type"] != "material"]) == 3
    assert len([row for row in remaining if row["item_type"] == "material"]) == 1
    snapshot_id = applied.json()["snapshot_id"]
    restored = client.post(f"/api/code-rule-snapshots/{snapshot_id}/restore?reason=规则恢复测试", headers=headers)
    assert restored.status_code == 200, restored.text
    editable_rule = client.get("/api/code-rules", headers=headers).json()[0]
    edited = client.patch(f"/api/code-rules/{editable_rule['id']}", headers=headers, json={
        "material_attribute": "受限编辑属性", "description": "只修改属性和说明",
        "reason": "验证编码规则受限编辑",
    })
    assert edited.status_code == 200, edited.text
    assert edited.json()["material_attribute"] == "受限编辑属性"
    forbidden = client.patch(f"/api/code-rules/{editable_rule['id']}", headers=headers, json={
        "material_attribute": "不应保存", "prefix": "88.8", "reason": "尝试越权修改前缀",
    })
    assert forbidden.status_code == 422
    event_id = client.get("/api/audit", headers=headers).json()[0]["id"]
    try:
        db.execute(text("UPDATE audit_events SET reason='tampered' WHERE id=:id"), {"id": event_id})
        db.commit()
        assert False, "审计记录不应允许修改"
    except DatabaseError:
        db.rollback()
