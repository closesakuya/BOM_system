import io
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from backend.app import models, schemas, services
from backend.app.material_import import complete_code


def test_prefix_allocation_and_rule_ambiguity(db):
    rule=models.CodeRule(item_type='material',prefix='10.',pattern='10.XXXX.0',active=True)
    db.add(rule);db.commit()
    assert complete_code(db,'10.01','material')[0]=='10.0100.0'
    assert complete_code(db,'10.01XX','material',reserved={'10.0100.0'})[0]=='10.0101.0'
    assert complete_code(db,'10.0105.0','material')[2] is False


def test_mixed_import_aliases_temporary_and_final_codes(client,headers,db):
    for kind,prefix,pattern in [('material','10.','10.XXXX.0'),('semi_finished','05.','05.XXX.01'),('unit','03.','03.XXX.01'),('machine','00.','00.XXX.01')]:
        db.add(models.CodeRule(item_type=kind,prefix=prefix,pattern=pattern,active=True))
    db.commit()
    wb=Workbook();ws=wb.active
    ws.append(['物料类型','物料编码','系统物料编码','名称','属性','单位','关键器件码','源于历史号'])
    ws.append(['原材料','10.01','10.01','测试导入件','外购','pcs','ab1.c','old-001'])
    ws.append(['半成品','05.7',None,'测试组件',None,'pcs',None,None])
    ws.append(['单元','03.7',None,'测试单元',None,'pcs',None,None])
    ws.append(['整机','00.7',None,'测试整机',None,'pcs',None,None])
    ws.append(['原材料','99.XXXX',None,'临时新件','外协','pcs',None,None])
    buf=io.BytesIO();wb.save(buf)
    response=client.post('/api/imports/materials/preview',headers=headers,files={'file':('test.xlsx',buf.getvalue())})
    assert response.status_code==200,response.text
    data=response.json();rows=data['rows']
    assert all(r['status']!='error' for r in rows),rows
    assert rows[0]['payload']['code']=='10.0100.0'
    assert rows[4]['payload']['is_formally_imported'] is False
    assert not rows[4]['payload'].get('historical_item_code')
    response=client.post(f"/api/imports/{data['batch']['id']}/commit",headers=headers,json={'row_ids':[r['id'] for r in rows],'similarity_confirmed_row_ids':[r['id'] for r in rows]})
    assert response.status_code==200,response.text
    assert len(response.json()['created'])==5
    item=db.scalar(select(models.Item).where(models.Item.code=='10.0100.0'))
    assert item.key_component_code=='AB1.C' and item.historical_item_code=='old-001'
    report=client.get(f"/api/imports/{data['batch']['id']}/report",headers=headers)
    assert report.status_code==200
    sheet=load_workbook(io.BytesIO(report.content)).active
    assert sheet.cell(2,2).value=='原材料'
    assert sheet.cell(2,3).value=='10.01' and sheet.cell(2,4).value=='10.0100.0'
    assert sheet.cell(2,11).value=='已导入'


def test_delete_frees_code_without_reusing_identity(db):
    from backend.app.item_lifecycle import delete_item
    actor=db.scalar(select(models.User))
    payload=schemas.ItemCreate(item_type='material',code='10.8100.0',name='误建',source_type='purchased',similarity_confirmed=True)
    item=services.create_item(db,payload,actor);old_id=item.id
    delete_item(db,item,actor,'清理误建')
    replacement=services.create_item(db,payload,actor)
    assert replacement.id>old_id
    assert db.get(models.Item,old_id) is None


def test_promoted_target_delete_keeps_source_snapshot_not_reused_link(db):
    from backend.app.item_lifecycle import delete_item
    actor=db.scalar(select(models.User))
    rule=models.CodeRule(item_type='material',prefix='10.',pattern='10.XXXX.0',active=True)
    db.add(rule);db.commit()
    source=services.create_item(db,schemas.ItemCreate(item_type='material',name='临时来源',source_type='purchased',is_formally_imported=False,similarity_confirmed=True),actor)
    target=services.promote_material(db,source,schemas.PromoteMaterialIn(code='10.8510.0',code_rule_id=rule.id,
        name='转正式目标',source_type='purchased',reason='转正式',similarity_confirmed=True),actor)
    target_id=target.id
    delete_item(db,target,actor,'删除误导入目标')
    trace=services.promotion_trace(db,source)
    assert trace['deleted_promoted_to']['code']=='10.8510.0'
    assert not trace.get('promoted_to')
    replacement=services.create_item(db,schemas.ItemCreate(code='10.8510.0',item_type='material',name='重新建档',source_type='purchased',similarity_confirmed=True),actor)
    assert replacement.id!=target_id
    assert not services.promotion_trace(db,replacement)['promoted_from']
    import pytest
    from fastapi import HTTPException
    with pytest.raises(HTTPException):delete_item(db,source,actor,'封存不可删')


def test_import_commit_reallocates_auto_and_atomic_conflicts(client,headers,db):
    actor=db.scalar(select(models.User))
    db.add(models.CodeRule(item_type='material',prefix='10.',pattern='10.XXXX.0',active=True));db.commit()
    def preview(codes):
        wb=Workbook();ws=wb.active;ws.append(['物料编码','名称','属性'])
        for n,code in enumerate(codes):ws.append([code,f'导入事务组件{n}','外购'])
        buf=io.BytesIO();wb.save(buf)
        response=client.post('/api/imports/materials/preview',headers=headers,files={'file':('atomic.xlsx',buf.getvalue())})
        assert response.status_code==200,response.text
        return response.json()
    def occupy(code):return services.create_item(db,schemas.ItemCreate(code=code,item_type='material',name=code,source_type='purchased',similarity_confirmed=True),actor)
    batch=preview(['10.02'])
    occupy('10.0200.0')
    response=client.post(f"/api/imports/{batch['batch']['id']}/commit",headers=headers,json={'row_ids':[r['id'] for r in batch['rows']],'similarity_confirmed_row_ids':[r['id'] for r in batch['rows']]})
    assert response.status_code==200,response.text
    assert response.json()['created'][0]['code']=='10.0201.0'
    batch=preview(['10.03','10.0400.0'])
    occupy('10.0400.0')
    response=client.post(f"/api/imports/{batch['batch']['id']}/commit",headers=headers,json={'row_ids':[r['id'] for r in batch['rows']],'similarity_confirmed_row_ids':[r['id'] for r in batch['rows']]})
    assert response.status_code==409
    assert db.scalar(select(models.Item).where(models.Item.code=='10.0300.0')) is None
    rows=client.get(f"/api/imports/{batch['batch']['id']}",headers=headers).json()['rows']
    assert any(r['status']=='error' for r in rows)
