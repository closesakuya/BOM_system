from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select

from . import models, services
from .config import PROJECT_ROOT, settings
from .database import SessionLocal
from .migration import LEGACY_PATH, RULE_PATH, preflight


def validate_initial_database() -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    migration: dict[str, Any] | None = None
    if LEGACY_PATH.is_file() and RULE_PATH.is_file():
        try:
            migration = preflight()
        except Exception as exc:  # Existing deployed databases do not depend on one-time source files.
            warnings.append(f"历史迁移源文件无法复核：{exc}")
    else:
        missing = [str(path.relative_to(PROJECT_ROOT)) for path in (LEGACY_PATH, RULE_PATH) if not path.is_file()]
        warnings.append(f"未携带一次性历史迁移源文件，跳过源数据复核：{', '.join(missing)}")
    with SessionLocal() as db:
        items = list(db.scalars(select(models.Item)))
        lines = list(db.scalars(select(models.BOMLine)))
        by_id = {item.id: item for item in items}
        counts = Counter(item.item_type for item in items)
        invalid_relations = []
        graph: dict[int, list[int]] = defaultdict(list)
        for line in lines:
            parent, child = by_id[line.parent_item_id], by_id[line.child_item_id]
            graph[parent.id].append(child.id)
            if child.item_type not in services.ALLOWED_CHILDREN.get(parent.item_type, set()):
                invalid_relations.append({"parent": parent.code, "child": child.code})
        if invalid_relations:
            errors.append(f"存在 {len(invalid_relations)} 条非法类型关系")

        visiting: set[int] = set()
        visited: set[int] = set()
        cycle = False

        def visit(node: int) -> None:
            nonlocal cycle
            if node in visiting:
                cycle = True
                return
            if node in visited:
                return
            visiting.add(node)
            for child in graph.get(node, []):
                visit(child)
            visiting.remove(node)
            visited.add(node)

        for item_id in by_id:
            visit(item_id)
        if cycle:
            errors.append("BOM 图存在环路")

        duplicate_official_codes = db.scalar(
            select(func.count()).select_from(models.Item).where(
                models.Item.is_formally_imported.is_(True),
                models.Item.code.in_(
                    select(models.Item.code)
                    .where(models.Item.is_formally_imported.is_(True))
                    .group_by(models.Item.code)
                    .having(func.count() > 1)
                ),
            )
        ) or 0
        if duplicate_official_codes:
            errors.append("存在重复正式系统物料编码")
        duplicate_official_versions = db.scalar(
            select(func.count()).select_from(models.Item).where(
                models.Item.is_formally_imported.is_(True),
                models.Item.id.in_(
                    select(func.min(models.Item.id))
                    .where(models.Item.is_formally_imported.is_(True))
                    .group_by(models.Item.version_series, models.Item.version_number)
                    .having(func.count() > 1)
                ),
            )
        ) or 0
        if duplicate_official_versions:
            errors.append("存在重复正式物料版本记录")
        unofficial_child_refs = db.scalar(
            select(func.count()).select_from(models.BOMLine)
            .join(models.Item, models.Item.id == models.BOMLine.child_item_id)
            .where(models.Item.is_formally_imported.is_(False))
        ) or 0
        unofficial_alternative_refs = db.scalar(
            select(func.count()).select_from(models.AlternativeMember)
            .join(models.Item, models.Item.id == models.AlternativeMember.item_id)
            .where(models.Item.is_formally_imported.is_(False))
        ) or 0
        invalid_unofficial_types = db.scalar(
            select(func.count()).select_from(models.Item).where(
                models.Item.is_formally_imported.is_(False),
                models.Item.item_type != "material",
            )
        ) or 0
        if unofficial_child_refs:
            errors.append("存在未正式导入原材料上级 BOM 引用")
        if unofficial_alternative_refs:
            errors.append("存在未正式导入原材料替代组引用")
        if invalid_unofficial_types:
            errors.append("存在非原材料被标记为未正式导入")
        unofficial_codes = [item.code for item in items if not item.is_formally_imported]
        duplicate_unofficial_codes = len(unofficial_codes) - len(set(unofficial_codes))
        invalid_unofficial_codes = sum(
            1 for code in unofficial_codes if re.fullmatch(r"99\.\d{4}\.\d", code) is None
        )
        formal_reserved_codes = sum(
            1 for item in items if item.is_formally_imported and re.fullmatch(r"99\.\d{4}\.\d", item.code)
        )
        invalid_unofficial_status = sum(
            1 for item in items
            if (not item.is_formally_imported and item.unofficial_status not in services.UNOFFICIAL_STATUSES)
            or (item.is_formally_imported and item.unofficial_status is not None)
        )
        invalid_key_component_codes = sum(
            1 for item in items
            if item.key_component_code and services.KEY_COMPONENT_CODE_PATTERN.fullmatch(item.key_component_code) is None
        )
        linked_sources = set(db.scalars(select(models.MaterialPromotionLink.source_item_id)))
        linked_sources.update(db.scalars(select(models.DeletedMaterialPromotion.source_item_id)))
        archived_without_link = sum(
            1 for item in items if item.unofficial_status == "archived" and item.id not in linked_sources
        )
        if duplicate_unofficial_codes:
            errors.append("存在重复未正式原材料编号")
        if invalid_unofficial_codes:
            errors.append("存在不符合 99.XXXX.X 的未正式原材料编号")
        if formal_reserved_codes:
            errors.append("存在正式物料占用 99 专属号段")
        if invalid_unofficial_status:
            errors.append("存在未正式状态与正式导入状态不一致")
        if invalid_key_component_codes:
            errors.append("存在格式无效的关键器件码")
        if archived_without_link:
            errors.append("存在没有转正式关联的封存来源")

        invalid_market_share_groups: list[int] = []
        for group_id in db.scalars(select(models.AlternativeGroup.id)):
            total = sum(
                (Decimal(member.market_share) for member in db.scalars(
                    select(models.AlternativeMember).where(
                        models.AlternativeMember.group_id == group_id,
                        models.AlternativeMember.active.is_(True),
                    )
                )),
                Decimal("0"),
            )
            if total != Decimal("100.00"):
                invalid_market_share_groups.append(int(group_id))
        if invalid_market_share_groups:
            errors.append(f"存在 {len(invalid_market_share_groups)} 个替代组市场占比合计不为 100%")

        sample_path = PROJECT_ROOT / "data" / "samples" / "00.005.01-TN水质分析仪-D60-YJ-BZ-A1-I-16C-11-0-生产BOM.xlsx"
        expected_headers = ["序号", "物料编码", "名称", "规格型号", "单位", "数量", "备注（黄色：修改后名称 ）", "上一版本名称", "发票名称", "材料属性", "变更记录"]
        header_match: bool | None = None
        if sample_path.is_file():
            try:
                workbook = load_workbook(sample_path, read_only=True, data_only=True)
                sheet = workbook.active
                sample_headers = [sheet.cell(2, column).value for column in range(1, 12)]
                header_match = sample_headers == expected_headers
                workbook.close()
                if not header_match:
                    errors.append("生产 BOM 样例字段与确认口径不一致")
            except Exception as exc:
                warnings.append(f"生产 BOM 样例无法复核：{exc}")
        else:
            warnings.append("未携带一次性生产 BOM 样例，跳过样例表头复核")
        root = db.scalar(select(models.Item).where(
            models.Item.code == "00.005.01",
            models.Item.is_formally_imported.is_(True),
        ))
        production_rows = len(services.production_bom(db, root.id)) if root else 0
        technical_rows = len(services.technical_bom(db, root.id)) if root else 0
        if not root or not production_rows or not technical_rows:
            errors.append("代表性整机无法展开 BOM")

        try:
            from .path_bom import EffectiveBOM
            EffectiveBOM(db).validate()
        except Exception as exc:
            errors.append(f'路径选配有效组成校验失败：{exc}')

        report = {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "database": str(Path(settings.database_url.removeprefix("sqlite:///"))),
            "passed": not errors,
            "errors": errors,
            "warnings": warnings,
            "items": len(items),
            "item_counts": dict(counts),
            "bom_lines": len(lines),
            "invalid_relations": invalid_relations,
            "cycle_detected": cycle,
            "duplicate_official_codes": duplicate_official_codes,
            "duplicate_official_versions": duplicate_official_versions,
            "unofficial_materials": counts.get("material", 0) and sum(
                1 for item in items if item.item_type == "material" and not item.is_formally_imported
            ),
            "unofficial_child_refs": unofficial_child_refs,
            "unofficial_alternative_refs": unofficial_alternative_refs,
            "invalid_unofficial_types": invalid_unofficial_types,
            "duplicate_unofficial_codes": duplicate_unofficial_codes,
            "invalid_unofficial_codes": invalid_unofficial_codes,
            "formal_reserved_codes": formal_reserved_codes,
            "invalid_unofficial_status": invalid_unofficial_status,
            "invalid_key_component_codes": invalid_key_component_codes,
            "archived_without_link": archived_without_link,
            "invalid_market_share_groups": invalid_market_share_groups,
            "migration_source_available": migration is not None,
            "migration_source_unique_items": migration["unique_items"] if migration else None,
            "migration_source_unique_bom_lines": migration["unique_bom_lines"] if migration else None,
            "resolved_cross_type_collisions": migration["resolved_material_machine_collisions"] if migration else None,
            "production_sample_headers_match": header_match,
            "representative_root": "00.005.01",
            "representative_technical_rows": technical_rows,
            "representative_production_rows": production_rows,
            "note": "生产 BOM 样例仅复核历史基础字段；当前 Excel 使用动态选配列，网页仍按行展示，不把样例作为运行输入。",
        }
    settings.report_dir.mkdir(parents=True, exist_ok=True)
    target = settings.report_dir / "v1_2-database-validation.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    result = validate_initial_database()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["passed"] else 1)
