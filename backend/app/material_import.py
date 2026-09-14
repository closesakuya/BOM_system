"""Four-type archive import with deterministic, non-reserving code previews."""
from __future__ import annotations
import io
import json
import re
from collections import Counter

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select, text

from . import models, schemas, services

TYPE_MAP = {'原材料':'material','半成品':'semi_finished','组合半成品':'semi_finished','单元':'unit','整机':'machine'}


def complete_code(db, value, kind, rule_id=None, reserved=None):
    value=str(value or '').strip().upper()
    reserved=set(reserved or [])
    if re.fullmatch(r'\d+\.\d+\.\d+',value):
        return value, rule_id, False
    rules=list(db.scalars(select(models.CodeRule).where(models.CodeRule.active.is_(True),models.CodeRule.item_type==kind)))
    candidates=[]
    for rule in rules:
        if rule_id and rule.id!=rule_id:
            continue
        template=services.code_rule_template(rule).upper()
        body,sep,version=template.rpartition('.')
        if not sep or not value:
            continue
        body_input=value
        if len(value)>len(body):
            if not value.startswith(body[:body.find('X')]) or value[len(body):]!=('.'+version):
                continue
            body_input=value[:len(body)]
        if len(body_input)>len(body):
            continue
        pattern=[]
        valid=True
        for i,ch in enumerate(body):
            entered=body_input[i] if i<len(body_input) else 'X'
            if ch=='X':
                if entered not in '0123456789X':valid=False;break
                pattern.append(entered)
            else:
                if i<len(body_input) and entered!=ch:valid=False;break
                pattern.append(ch)
        if valid:
            candidates.append((rule,''.join(pattern),version.replace('X','0')))
    if len(candidates)!=1:
        raise HTTPException(422,'前缀匹配多条编码规则，请选择大类、小类和编码规则' if candidates else '没有匹配的编码规则，请修正前缀或选择规则')
    rule,body,version=candidates[0]
    indexes=[i for i,c in enumerate(body) if c=='X']
    if len(indexes)>6:
        raise HTTPException(422,'编码范围过大，请补充前缀')
    occupied=set(db.scalars(select(models.Item.code)))|reserved
    for number in range(10**len(indexes)):
        chars=list(body)
        for index,digit in zip(indexes,str(number).zfill(len(indexes))):chars[index]=digit
        code=''.join(chars)+'.'+version
        if code not in occupied:
            return code,rule.id,True
    raise HTTPException(409,'输入前缀范围内的编号已用完，请修改前缀')


def prepare(db,data,reserved,unofficial_offset):
    from .excel_io import validate_material_import_payload
    errors=list(data.get('alias_errors') or [])
    kind=data.get('item_type','material')
    original=str(data.get('code_input') if data.get('code_input') is not None else data.get('code') or '').strip()
    unofficial=original=='99' or original.startswith('99.') or not data.get('is_formally_imported',True)
    if kind not in services.ITEM_TYPES:errors.append('物料类型无效')
    if unofficial and kind!='material':errors.append('99 号段及未正式状态仅允许用于原材料')
    data['is_formally_imported']=not unofficial
    if unofficial:
        data['code']=services.preview_unofficial_code(db,unofficial_offset)
        data['code_rule_id']=None
        if not data.get('historical_item_code') and original and original not in {'99','99.'} and 'X' not in original.upper():
            data['historical_item_code']=original
    else:
        try:
            code,rule,auto=complete_code(db,original,kind,data.get('code_rule_id'),reserved)
            data.update(code=code,code_rule_id=rule,auto_code=auto,code_input=original)
            if code in reserved:errors.append('与本批其他行指定编码重复')
        except HTTPException as exc:errors.append(str(exc.detail))
    if kind!='material':
        data['source_type']=None
        if data.get('key_component_code'):errors.append('关键器件码仅适用于原材料')
    data['status']=data.get('status') or ('active' if kind=='material' else 'trial')
    try: services.validate_item_business(kind,data.get('source_type'),data.get('status'))
    except HTTPException as exc:errors.append(str(exc.detail))
    if data.get('machine_model') and kind!='machine':errors.append('机型仅适用于整机')
    status,message,similarity=validate_material_import_payload(db,data)
    if message and status=='error':errors.append(message)
    if data.get('code') and not unofficial:reserved.add(data['code'])
    notice='按 99 号段识别为未正式导入，提交时自动分配编号' if original=='99' or original.startswith('99.') else ''
    return ('error' if errors else status), '；'.join(errors) if errors else ('；'.join(filter(None,[notice,message])) or None),similarity


def revalidate(db,batch):
    rows=list(db.scalars(select(models.ImportRow).where(models.ImportRow.batch_id==batch.id).order_by(models.ImportRow.row_number)))
    # Explicit codes take precedence over automatically filled rows regardless of row order.
    reserved={str(d.get('code_input') or d.get('code')) for r in rows if r.status not in {'cancelled','imported'}
              for d in [json.loads(r.payload_json)] if re.fullmatch(r'\d+\.\d+\.\d+',str(d.get('code_input') or d.get('code') or '')) and not d.get('auto_code') and d.get('is_formally_imported',True)}
    seen_explicit=set();offset=0
    for row in rows:
        if row.status in {'cancelled','imported'}:continue
        data=json.loads(row.payload_json)
        inp=str(data.get('code_input') or data.get('code') or '')
        explicit=inp in reserved and not data.get('auto_code')
        if explicit and inp not in seen_explicit:reserved.discard(inp)
        status,message,similarity=prepare(db,data,reserved,offset)
        if explicit:seen_explicit.add(inp)
        if not data.get('is_formally_imported',True):offset+=1
        row.payload_json=services.dumps(data);row.status=status;row.error_message=message;row.similarity_json=services.dumps(similarity)
    db.flush()
    batch.stats_json=services.dumps(dict(Counter(row.status for row in rows)))
    return rows


async def preview(db,upload,actor):
    from .excel_io import HEADER_ALIASES,SOURCE_MAP,import_row_dict
    content=await upload.read()
    if len(content)>20*1024*1024:raise HTTPException(413,'文件不能超过 20MB')
    try:wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    except Exception as exc:raise HTTPException(422,'无法读取 Excel') from exc
    ws=wb[wb.sheetnames[0]]
    aliases={**HEADER_ALIASES,'item_type':{'物料类型','类型'},'historical_item_code':{'源于历史号','历史物料号','原物料识别码'},'status':{'状态'},'machine_model':{'机型'}}
    mapping={};header=0
    for index,values in enumerate(ws.iter_rows(max_row=10,values_only=True),1):
        mapping={field:[i for i,v in enumerate(values) if str(v or '').strip() in names] for field,names in aliases.items()}
        if mapping['name'] and mapping['code']:header=index;break
    if not header:wb.close();raise HTTPException(422,'无法识别表头，需要名称及物料编码列')
    batch=models.ImportBatch(filename=upload.filename or '物料导入.xlsx',created_by=actor.id)
    db.add(batch);db.flush()
    for index,values in enumerate(ws.iter_rows(min_row=header+1,values_only=True),header+1):
        data={};errors=[]
        for field,columns in mapping.items():
            nonempty=list(dict.fromkeys(str(values[c]).strip() for c in columns if c<len(values) and values[c] is not None and str(values[c]).strip()))
            if len(nonempty)>1:errors.append(f'{field} 同义列存在不同值，请编辑修正')
            data[field]=nonempty[0] if nonempty else None
        if not data.get('name') and not data.get('code'):continue
        kind=data.get('item_type') or '原材料'
        data['item_type']=TYPE_MAP.get(kind,kind)
        data['source_type']=SOURCE_MAP.get(data.get('source_type'))
        state=data.pop('import_status',None)
        data['is_formally_imported']=state!='未正式导入'
        data['import_status_error']=None if state in {None,'','正式导入','未正式导入'} else '导入状态无效'
        data['status']={'在用':'active','试制':'trial','停用':'disabled'}.get(data.get('status'),data.get('status'))
        data['unit']=data.get('unit') or 'pcs'
        data['code_input']=data.get('code');data['alias_errors']=errors
        db.add(models.ImportRow(batch_id=batch.id,row_number=index,payload_json=services.dumps(data),status='ready'))
    wb.close();db.flush()
    rows=revalidate(db,batch)
    services.audit(db,actor.id,'preview','import_batch',batch.id,'预览物料导入',after={'filename':batch.filename})
    db.commit()
    return {'batch':{'id':batch.id,'filename':batch.filename,'status':batch.status,'stats':json.loads(batch.stats_json)},'rows':[import_row_dict(r) for r in rows]}


def update_row(db,batch,row,payload,actor):
    from .excel_io import import_row_dict
    if batch.status=='completed':raise HTTPException(409,'已完成的导入不可编辑')
    before=import_row_dict(row)
    if payload.cancelled:row.status='cancelled'
    else:
        data=payload.model_dump(exclude={'cancelled'})
        prior=json.loads(row.payload_json)
        # Keep prefix intention only if the user did not change the displayed code.
        if data.get('code')!=prior.get('code'):
            data['code_input']=data.get('code')
        else:data['code_input']=prior.get('code_input') or data.get('code')
        data['alias_errors']=[];data['import_status_error']=None
        row.payload_json=services.dumps(data);row.status='ready'
    db.flush();revalidate(db,batch)
    services.audit(db,actor.id,'update','import_row',row.id,'编辑物料导入预览',before,import_row_dict(row))
    db.commit();return import_row_dict(row)


def commit(db,batch,payload,actor):
    if batch.status=='completed':raise HTTPException(409,'本批已经提交，请勿重复导入')
    # Serialize reservations against concurrent creates before final revalidation.
    db.execute(text("UPDATE code_sequences SET next_value=next_value WHERE name='unofficial_material'"))
    rows=[r for r in revalidate(db,batch) if r.status not in {'cancelled','imported'}]
    errors=[r for r in rows if r.status=='error']
    if errors:
        db.commit()
        raise HTTPException(409,'请修正或取消错误行：'+'；'.join(f'{r.row_number}: {r.error_message}' for r in errors))
    if set(payload.row_ids)!={r.id for r in rows}:raise HTTPException(409,'请提交全部未取消行')
    confirmed=set(payload.similarity_confirmed_row_ids)
    if any(r.status=='red' and r.id not in confirmed for r in rows):raise HTTPException(409,'请确认高度相似物料')
    result=[]
    try:
        for row in rows:
            data=json.loads(row.payload_json)
            fields={k:v for k,v in data.items() if k in schemas.ItemCreate.model_fields}
            fields.update(similarity_confirmed=row.id in confirmed,reason=payload.reason)
            item=services.create_item(db,schemas.ItemCreate(**fields),actor,commit=False,historical_item_code=data.get('historical_item_code'))
            data['code']=item.code;row.payload_json=services.dumps(data);row.status='imported';row.selected=True;row.decision='create'
            result.append(services.item_dict(item))
        batch.status='completed';batch.stats_json=services.dumps(dict(Counter(r.status for r in rows)))
        services.audit(db,actor.id,'commit','import_batch',batch.id,payload.reason,after={'created':result})
        db.commit()
    except Exception:db.rollback();raise
    return {'created':result,'batch_status':'completed'}
