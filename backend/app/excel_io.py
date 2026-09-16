from __future__ import annotations

import io
import json
import re
import zipfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import models, schemas, services


TECHNICAL_HEADERS = [
    "序号", "层级", "物料编码", "名称", "规格型号", "单位", "数量", "备注",
    "上一版本名称", "发票名称", "材料属性", "选配状态", "替代组", "市场占比", "变更记录",
]
PRODUCTION_HEADERS = [
    "序号", "物料编码", "名称", "规格型号", "单位", "数量", "备注",
    "上一版本名称", "发票名称", "材料属性", "选配状态", "替代组", "市场占比", "变更记录",
]


def safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", value).strip().rstrip(".")
    return cleaned[:160] or "BOM"


def style_sheet(ws, column_count: int) -> None:  # type: ignore[no-untyped-def]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(column_count)}{max(ws.max_row, 2)}"
    widths = [8, 18, 22, 32, 42, 10, 12, 38, 28, 28, 24, 16, 24, 14, 38]
    for index in range(1, column_count + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    for row in ws.iter_rows(min_row=3):
        ws.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        if any("停用" in str(cell.value or "") for cell in row):
            for cell in row:
                cell.fill = warning_fill
        if any("备用选配" in str(cell.value or "") for cell in row):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="DDEBF7")


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def technical_level_labels(rows: list[dict[str, Any]]) -> list[str]:
    """Render flattened BOM paths as a readable tree in the Excel level column."""
    last_child_by_parent: dict[tuple[int, ...], tuple[int, ...]] = {}
    for row in rows:
        path = tuple(row["path"])
        last_child_by_parent[path[:-1]] = path

    labels: list[str] = []
    for row in rows:
        path = tuple(row["path"])
        level = int(row["level"])
        prefix: list[str] = []
        for ancestor_level in range(1, level):
            ancestor_path = path[:ancestor_level + 1]
            ancestor_is_last = last_child_by_parent.get(ancestor_path[:-1]) == ancestor_path
            prefix.append("   " if ancestor_is_last else "│  ")
        is_last = last_child_by_parent.get(path[:-1]) == path
        labels.append(f"{''.join(prefix)}{'└─' if is_last else '├─'} {level}级")
    return labels


def technical_workbook(
    db: Session, root_id: int, *, show_alternatives: bool = True, expand_materials: bool | None = None,
) -> tuple[str, bytes]:
    from .bom_excel import export_workbook
    return export_workbook(db, root_id, show_alternatives=show_alternatives, expand_materials=expand_materials)


def production_workbook(
    db: Session, root_id: int, *, show_alternatives: bool = True, expand_materials: bool | None = None,
) -> tuple[str, bytes]:
    from .bom_excel import export_workbook
    return export_workbook(db, root_id, production=True, show_alternatives=show_alternatives, expand_materials=expand_materials)


def latest_item_change_reasons(db: Session, item_ids: set[int]) -> dict[int, str]:
    if not item_ids:
        return {}
    result: dict[int, str] = {}
    events = db.scalars(
        select(models.AuditEvent)
        .where(
            models.AuditEvent.entity_type == "item",
            models.AuditEvent.entity_id.in_(item_ids),
        )
        .order_by(models.AuditEvent.created_at.desc(), models.AuditEvent.id.desc())
    )
    for event in events:
        if event.entity_id is not None and event.entity_id not in result:
            result[event.entity_id] = event.reason
    return result


def batch_bom_zip(
    db: Session, item_ids: list[int], bom_type: str, *,
    show_alternatives: bool = True, expand_materials: bool | None = None,
) -> tuple[str, bytes]:
    if not item_ids:
        raise HTTPException(status_code=422, detail="至少选择一个对象")
    output = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for item_id in item_ids:
            filename, content = (
                technical_workbook(db, item_id, show_alternatives=show_alternatives, expand_materials=expand_materials)
                if bom_type == "technical"
                else production_workbook(db, item_id, show_alternatives=show_alternatives, expand_materials=expand_materials)
            )
            base = filename
            index = 2
            while filename in used:
                filename = base.replace(".xlsx", f"-{index}.xlsx")
                index += 1
            used.add(filename)
            archive.writestr(filename, content)
    return f"批量{('技术' if bom_type == 'technical' else '生产')}BOM.zip", output.getvalue()


HEADER_ALIASES = {
    "code": {"物料识别码", "系统物料编码", "物料编码"},
    "auxiliary_code": {"辅助索引码", "旧物料编码"},
    "name": {"名称", "*名称", "物料名称"},
    "specification": {"规格型号", "*规格型号", "规格"},
    "source_type": {"属性", "*属性", "来源属性"},
    "unit": {"单位", "*单位"},
    "remark": {"备注"},
    "previous_version_name": {"上一版本名称"},
    "invoice_name": {"发票名称"},
    "material_attribute": {"材料属性", "*材料属性"},
    "key_component_code": {"关键器件码"},
    "import_status": {"导入状态"},
}
SOURCE_MAP = {"外购": "purchased", "外协": "outsourced", "自制": "self_made"}
SOURCE_LABELS = {"purchased": "外购", "outsourced": "外协", "self_made": "自制"}
STATUS_LABELS = {"active": "在用", "trial": "试制", "disabled": "停用"}
TYPE_LABELS = {"material": "原材料", "semi_finished": "半成品", "unit": "单元", "machine": "整机"}


def item_basics_workbook(
    db: Session,
    items: list[models.Item],
    item_type: str,
    *,
    selected: bool,
) -> tuple[str, bytes]:
    return _item_basics_workbook(db, items, item_type, selected=selected)


def material_basics_workbook(db: Session, items: list[models.Item], *, selected: bool):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '原材料基本信息'
    sheet.append(['序号','关键器件码','物料编码','*名称','*规格型号','*属性','*单位','备注','上一版本名称','发票名称','*材料属性','变更记录','当前在用归属产品机型'])
    requested={item.id for item in items}
    history=defaultdict(list)
    labels={'name':'名称','specification':'规格型号','remark':'备注','status':'状态','source_type':'来源',
            'quantity':'数量','line_remark':'组成备注','unit':'单位','key_component_code':'关键器件码',
            'material_attribute':'材料属性','invoice_name':'发票名称','previous_version_name':'上一版本名称'}
    for event in db.scalars(select(models.AuditEvent).order_by(models.AuditEvent.created_at.desc(),models.AuditEvent.id.desc())):
        before=json.loads(event.before_json) if event.before_json else {}
        after=json.loads(event.after_json) if event.after_json else {}
        before=before if isinstance(before,dict) else {}
        after=after if isinstance(after,dict) else {}
        subject=event.entity_id if event.entity_type=='item' else (after.get('parent_item_id') or before.get('parent_item_id')) if event.entity_type=='bom_line' else None
        if subject not in requested:continue
        changes=[]
        for field,label in labels.items():
          if before.get(field)!=after.get(field):
                values = SOURCE_LABELS if field=='source_type' else {'active':'在用','disabled':'停用','trial':'试制'} if field=='status' else {}
                changes.append(f'{label}：{values.get(before.get(field),before.get(field)) or "—"} → {values.get(after.get(field),after.get(field)) or "—"}')
        if event.entity_type=='bom_line':
            old=before.get('child') or {};new=after.get('child') or {}
            changes.insert(0,f'组成：{old.get("code", "—")} {old.get("name", "")} → {new.get("code", "—")} {new.get("name", "")}')
        history[subject].append(f'{event.created_at:%Y-%m-%d %H:%M} '+('；'.join(changes) or '档案操作')+f'；说明：{event.reason}')
    refs=services.model_references_for_items(db,requested)
    for index,item in enumerate(items,1):
        record='\n'.join(history[item.id])
        if len(record)>32000:
            detail=workbook.create_sheet(f'历史-{index}')
            detail.append(['物料编码','变更记录'])
            for event_text in history[item.id]:detail.append([item.code,event_text[:32767]])
            record=record[:31000]+f'\n完整记录见工作表 历史-{index}'
        sheet.append([index,item.key_component_code,item.code,item.name,item.specification,SOURCE_LABELS.get(item.source_type,''),
                      item.unit,item.remark,item.previous_version_name,item.invoice_name,item.material_attribute,record,
                      '、'.join(refs[item.id]['current'])])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78')
    for index,width in enumerate([8,16,22,38,38,10,10,36,30,26,24,64,32],1):sheet.column_dimensions[get_column_letter(index)].width=width
    return f'原材料基本信息-{"已选" if selected else "筛选结果"}.xlsx',workbook_bytes(workbook)


def _item_basics_workbook(
    db: Session,
    items: list[models.Item],
    item_type: str,
    *,
    selected: bool,
) -> tuple[str, bytes]:
    if item_type == 'material':
        return material_basics_workbook(db, items, selected=selected)
    common = ["系统物料编码", "名称", "规格型号", "单位", "类型", "正式导入状态", "备注"]
    extra = {
        "material": [
            "来源", "历史物料号", "关键器件码", "封存状态", "来源未正式物料号",
            "转正式物料号", "材料属性", "发票名称", "上一版本名称",
        ],
        "semi_finished": ["状态", "组合半成品", "引用机型"],
        "unit": ["状态", "引用机型"],
        "machine": ["状态", "机型"],
    }[item_type]
    references = services.model_references_for_items(db, [item.id for item in items])
    promotion_links = list(db.scalars(select(models.MaterialPromotionLink))) if item_type == "material" else []
    sources_by_target: dict[int, list[str]] = defaultdict(list)
    target_by_source: dict[int, str] = {}
    for link in promotion_links:
        source = db.get(models.Item, link.source_item_id)
        target = db.get(models.Item, link.target_item_id)
        if source and target:
            sources_by_target[target.id].append(source.code)
            target_by_source[source.id] = target.code
    combination_ids: set[int] = set()
    if item_type == "semi_finished" and items:
        combination_ids = {
            int(parent_id)
            for parent_id in db.scalars(
                select(models.BOMLine.parent_item_id)
                .join(models.Item, models.Item.id == models.BOMLine.child_item_id)
                .where(
                    models.BOMLine.parent_item_id.in_([item.id for item in items]),
                    models.Item.item_type == "semi_finished",
                )
                .distinct()
            )
        }
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{TYPE_LABELS[item_type]}基本信息"
    headers = [*common, *extra]
    sheet.append(headers)
    for item in items:
        row: list[Any] = [
            services.item_dict(item)["code"],
            item.name,
            item.specification,
            item.unit,
            TYPE_LABELS[item.item_type],
            "正式导入" if item.is_formally_imported else "未正式导入",
            item.remark,
        ]
        if item_type == "material":
            row.extend(
                [
                    SOURCE_LABELS.get(item.source_type or "", item.source_type),
                    item.historical_item_code,
                    item.key_component_code,
                    "已转正式（封存）" if item.unofficial_status == "archived" else "待转正式" if item.unofficial_status == "pending" else "—",
                    "、".join(sources_by_target.get(item.id, [])),
                    target_by_source.get(item.id),
                    item.material_attribute,
                    item.invoice_name,
                    item.previous_version_name,
                ]
            )
        elif item_type == "semi_finished":
            row.extend(
                [
                    STATUS_LABELS.get(item.status or "", item.status),
                    "组合半成品" if item.id in combination_ids else "普通半成品",
                    "、".join(references[item.id]["current"]),
                ]
            )
        elif item_type == "unit":
            row.extend(
                [
                    STATUS_LABELS.get(item.status or "", item.status),
                    "、".join(references[item.id]["current"]),
                ]
            )
        else:
            row.extend([STATUS_LABELS.get(item.status or "", item.status), item.machine_model])
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    widths = [22, 30, 36, 10, 12, 14, 36, 16, 22, 18, 20, 24, 24, 24, 24, 24]
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    scope = "已选" if selected else "筛选结果"
    filename = safe_filename(f"{TYPE_LABELS[item_type]}基本信息-{scope}.xlsx")
    return filename, workbook_bytes(workbook)


async def preview_code_rule_import(db: Session, upload: UploadFile) -> dict[str, Any]:
    content = await upload.read()
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取编码规则 Excel") from exc
    sheet = next((row for row in workbook.worksheets if "编码规则" in row.title), None)
    if not sheet:
        workbook.close()
        raise HTTPException(status_code=422, detail="未找到原材料分类及编码规则表")
    rules: list[dict[str, Any]] = []
    large_category = None
    seen: set[tuple[str, str]] = set()
    for row in sheet.iter_rows(min_row=3, values_only=True):
        large_category = str(row[1]).strip() if row[1] else large_category
        prefix = str(row[4]).strip() if row[4] else None
        pattern = str(row[5]).strip() if row[5] else None
        if not prefix or not pattern or prefix == "/" or pattern == "/":
            continue
        key = (prefix, pattern)
        if key in seen:
            continue
        seen.add(key)
        rules.append({
            "item_type": "material", "large_category": large_category,
            "small_category": str(row[2]).strip() if row[2] else None,
            "material_attribute": str(row[3]).strip() if row[3] else None,
            "prefix": prefix, "pattern": pattern,
            "description": str(row[6]).strip() if row[6] else None,
            "active": True, "sort_order": len(rules) + 1,
        })
    workbook.close()
    imported_types = {row["item_type"] for row in rules}
    existing = list(db.scalars(
        select(models.CodeRule).where(models.CodeRule.item_type.in_(imported_types))
    )) if imported_types else []
    fields = list(schemas.CodeRuleIn.model_fields)
    existing_rows = [{field: getattr(row, field) for field in fields} for row in existing]
    identity = lambda row: (row.get("item_type"), row.get("large_category"), row.get("small_category"), row.get("material_attribute"))
    old_map = {identity(row): row for row in existing_rows}
    new_map = {identity(row): row for row in rules}
    added = [row for key, row in new_map.items() if key not in old_map]
    removed = [row for key, row in old_map.items() if key not in new_map]
    changed = [{"before": old_map[key], "after": row} for key, row in new_map.items() if key in old_map and old_map[key] != row]
    return {"rules": rules, "diff": {"added": added, "changed": changed, "removed": removed}, "counts": {"rules": len(rules), "added": len(added), "changed": len(changed), "removed": len(removed)}}


def locate_header(ws) -> tuple[int, dict[str, int]]:  # type: ignore[no-untyped-def]
    for row_number in range(1, min(ws.max_row, 10) + 1):
        values = [str(ws.cell(row_number, col).value or "").strip() for col in range(1, ws.max_column + 1)]
        mapping: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for col, value in enumerate(values, 1):
                if value in aliases:
                    mapping[field] = col
                    break
        if "name" in mapping and "source_type" in mapping and "unit" in mapping:
            return row_number, mapping
    raise HTTPException(status_code=422, detail="无法识别原材料导入表头")


def recommend_code_for_row(db: Session, material_attribute: str | None) -> str | None:
    if not material_attribute:
        return None
    rule = db.scalar(
        select(models.CodeRule)
        .where(
            models.CodeRule.item_type == "material",
            models.CodeRule.active.is_(True),
            models.CodeRule.material_attribute.contains(material_attribute),
        )
        .order_by(models.CodeRule.sort_order, models.CodeRule.id)
    )
    if not rule:
        return None
    return services.next_code(db, rule.prefix, rule.pattern)["recommended_code"]


def validate_material_import_payload(
    db: Session,
    payload: dict[str, Any],
    *,
    batch_id: int | None = None,
    row_id: int | None = None,
) -> tuple[str, str | None, list[dict[str, Any]]]:
    errors: list[str] = []
    notices: list[str] = []
    name = str(payload.get("name") or "").strip()
    code = str(payload.get("code") or "").strip()
    formal = bool(payload.get("is_formally_imported", True))
    if payload.get("import_status_error"):
        errors.append(str(payload["import_status_error"]))
    if not name:
        errors.append("名称不能为空")
    if payload.get("item_type", "material") == "material" and payload.get("source_type") not in services.SOURCE_TYPES:
        errors.append("属性必须为外购、外协或自制")
    if not code:
        errors.append("系统物料编码不能为空")
    else:
        availability = services.code_availability(
            db,
            code,
            payload.get("code_rule_id"),
            payload.get("item_type", "material"),
            formal,
        )
        if not availability["valid"]:
            errors.append(availability["message"])
        elif formal and availability["exists"]:
            errors.append("正式系统物料编码已存在")
    try:
        normalized_key_code = services.normalize_key_component_code(payload.get("key_component_code"))
        payload["key_component_code"] = normalized_key_code
        if normalized_key_code:
            duplicates = list(db.scalars(
                select(models.Item).where(models.Item.key_component_code == normalized_key_code).limit(5)
            ))
            if duplicates:
                notices.append("关键器件码重复关联：" + "、".join(f"{row.code}｜{row.name}" for row in duplicates))
    except HTTPException as exc:
        errors.append(str(exc.detail))
    if formal and code and batch_id is not None:
        for candidate in db.scalars(
            select(models.ImportRow).where(
                models.ImportRow.batch_id == batch_id,
                models.ImportRow.id != row_id if row_id is not None else models.ImportRow.id.is_not(None),
                models.ImportRow.status != "cancelled",
            )
        ):
            other = json.loads(candidate.payload_json)
            if bool(other.get("is_formally_imported", True)) and other.get("code") == code:
                errors.append(f"与本批次第 {candidate.row_number} 行正式编码重复")
                break
    similarity = services.similarity_candidates(
        db,
        name,
        payload.get("specification"),
        payload.get("item_type", "material"),
    ) if name else []
    if errors:
        status = "error"
    elif any(row["level"] == "red" for row in similarity):
        status = "red"
    elif similarity:
        status = "yellow"
    else:
        status = "ready"
    return status, "；".join([*errors, *notices]) or None, similarity


def refresh_import_batch_stats(db: Session, batch: models.ImportBatch) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for status in db.scalars(
        select(models.ImportRow.status).where(models.ImportRow.batch_id == batch.id)
    ):
        counts[str(status)] += 1
    result = dict(counts)
    batch.stats_json = services.dumps(result)
    return result


def update_import_row(
    db: Session,
    batch: models.ImportBatch,
    row: models.ImportRow,
    payload: schemas.ImportRowEditIn,
    actor: models.User,
) -> dict[str, Any]:
    before = import_row_dict(row)
    if payload.cancelled:
        row.status = "cancelled"
        row.selected = False
        row.decision = "cancel"
        row.error_message = "用户取消本行，不参与导入"
        row.similarity_json = "[]"
    else:
        data = payload.model_dump(exclude={"cancelled"})
        data["item_type"] = "material"
        data["import_status_error"] = None
        data["code"] = str(data.get("code") or "").strip() or None
        data["name"] = str(data.get("name") or "").strip() or None
        if not data.get("is_formally_imported", True):
            candidate_code = str(data.get("code") or "")
            if candidate_code and not services.UNOFFICIAL_CODE_PATTERN.fullmatch(candidate_code):
                data["historical_item_code"] = candidate_code
            prior_unofficial = 0
            for prior in db.scalars(
                select(models.ImportRow).where(
                    models.ImportRow.batch_id == batch.id,
                    models.ImportRow.row_number < row.row_number,
                    models.ImportRow.status != "cancelled",
                )
            ):
                if not bool(json.loads(prior.payload_json).get("is_formally_imported", True)):
                    prior_unofficial += 1
            data["code"] = services.preview_unofficial_code(db, prior_unofficial)
        status, message, similarity = validate_material_import_payload(
            db,
            data,
            batch_id=batch.id,
            row_id=row.id,
        )
        row.payload_json = services.dumps(data)
        row.status = status
        row.error_message = message
        row.similarity_json = services.dumps(similarity)
        row.decision = "edit"
    refresh_import_batch_stats(db, batch)
    db.flush()
    services.audit(
        db,
        actor.id,
        "update",
        "import_row",
        row.id,
        "编辑原材料导入预览行" if not payload.cancelled else "取消原材料导入预览行",
        before,
        import_row_dict(row),
        batch_key=f"import-{batch.id}",
    )
    db.commit()
    return import_row_dict(row)


async def preview_material_import(db: Session, upload: UploadFile, actor: models.User) -> dict[str, Any]:
    content = await upload.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="导入文件不能超过 20MB")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取 Excel 文件") from exc
    ws = wb[wb.sheetnames[0]]
    header_row, mapping = locate_header(ws)
    batch = models.ImportBatch(filename=upload.filename or "原材料导入.xlsx", created_by=actor.id)
    db.add(batch)
    db.flush()
    counts = defaultdict(int)
    result_rows = []
    unofficial_offset = 0
    for row_number in range(header_row + 1, ws.max_row + 1):
        def value(field: str) -> Any:
            col = mapping.get(field)
            return ws.cell(row_number, col).value if col else None

        raw_code = str(value("code") or "").strip()
        name = str(value("name") or "").strip()
        if not name and not raw_code:
            continue
        source_label = str(value("source_type") or "").strip()
        status_label = str(value("import_status") or "").strip()
        formal = status_label in {"", "正式导入"}
        status_error = None if status_label in {"", "正式导入", "未正式导入"} else f"导入状态“{status_label}”无效，只能填写正式导入或未正式导入"
        payload = {
            "item_type": "material",
            "code": raw_code or None,
            "auxiliary_code": str(value("auxiliary_code") or "").strip() or None,
            "name": name,
            "specification": str(value("specification") or "").strip() or None,
            "source_type": SOURCE_MAP.get(source_label),
            "unit": str(value("unit") or "").strip() or "pcs",
            "remark": str(value("remark") or "").strip() or None,
            "previous_version_name": str(value("previous_version_name") or "").strip() or None,
            "invoice_name": str(value("invoice_name") or "").strip() or None,
            "material_attribute": str(value("material_attribute") or "").strip() or None,
            "key_component_code": str(value("key_component_code") or "").strip() or None,
            "historical_item_code": raw_code or None if not formal else None,
            "is_formally_imported": formal,
            "import_status_error": status_error,
        }
        if not formal:
            payload["code"] = services.preview_unofficial_code(db, unofficial_offset)
            unofficial_offset += 1
        elif not payload["code"]:
            payload["code"] = recommend_code_for_row(db, payload["material_attribute"])
        status, error_message, similarity = validate_material_import_payload(
            db,
            payload,
            batch_id=batch.id,
        )
        counts[status] += 1
        record = models.ImportRow(
            batch_id=batch.id,
            row_number=row_number,
            payload_json=services.dumps(payload),
            status=status,
            similarity_json=services.dumps(similarity),
            error_message=error_message,
        )
        db.add(record)
        db.flush()
        result_rows.append(import_row_dict(record))
    wb.close()
    batch.stats_json = services.dumps(dict(counts))
    services.audit(db, actor.id, "preview", "import_batch", batch.id, "预览原材料导入", after={"filename": batch.filename, "stats": dict(counts)})
    db.commit()
    return {"batch": {"id": batch.id, "filename": batch.filename, "status": batch.status, "stats": dict(counts)}, "rows": result_rows}


def import_row_dict(row: models.ImportRow) -> dict[str, Any]:
    return {
        "id": row.id,
        "batch_id": row.batch_id,
        "row_number": row.row_number,
        "payload": json.loads(row.payload_json),
        "status": row.status,
        "similarity": json.loads(row.similarity_json),
        "error_message": row.error_message,
        "selected": row.selected,
        "decision": row.decision,
    }


def commit_material_import(
    db: Session,
    batch: models.ImportBatch,
    payload: schemas.ImportCommitIn,
    actor: models.User,
) -> dict[str, Any]:
    all_rows = list(
        db.scalars(
            select(models.ImportRow)
            .where(models.ImportRow.batch_id == batch.id)
            .order_by(models.ImportRow.row_number)
        )
    )
    active_rows = [row for row in all_rows if row.status not in {"cancelled", "imported"}]
    if any(row.status == "error" for row in active_rows):
        invalid = [str(row.row_number) for row in active_rows if row.status == "error"]
        raise HTTPException(
            status_code=409,
            detail=f"第 {', '.join(invalid)} 行仍有错误；请修正或取消后再整批提交",
        )
    if set(payload.row_ids) != {row.id for row in active_rows}:
        raise HTTPException(status_code=409, detail="必须提交全部未取消行；不导入的行请先点击取消此行")
    rows = list(
        db.scalars(
            select(models.ImportRow).where(
                models.ImportRow.batch_id == batch.id,
                models.ImportRow.id.in_(payload.row_ids),
            )
        )
    )
    if len(rows) != len(set(payload.row_ids)):
        raise HTTPException(status_code=404, detail="部分导入行不存在")
    confirmed = set(payload.similarity_confirmed_row_ids)
    for row in rows:
        if row.status == "red" and row.id not in confirmed:
            raise HTTPException(status_code=409, detail=f"第 {row.row_number} 行需要确认高度相似物料")
    batch_key = services.new_batch_key()
    created = []
    try:
        for row in rows:
            data = json.loads(row.payload_json)
            data.pop("large_category", None)
            data.pop("small_category", None)
            data.pop("import_status_error", None)
            historical_item_code = data.pop("historical_item_code", None)
            data.update({"similarity_confirmed": row.id in confirmed, "reason": payload.reason, "components": []})
            item = services.create_item(
                db,
                schemas.ItemCreate(**data),
                actor,
                commit=False,
                historical_item_code=historical_item_code,
            )
            row.status = "imported"
            row.selected = True
            row.decision = "create"
            created.append(services.item_dict(item))
        batch.status = "completed"
        services.audit(db, actor.id, "commit", "import_batch", batch.id, payload.reason, after={"created": len(created)}, batch_key=batch_key)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"created": created, "batch_status": batch.status, "batch_key": batch_key}


def import_error_report(db: Session, batch_id: int) -> bytes:
    rows = list(db.scalars(select(models.ImportRow).where(models.ImportRow.batch_id == batch_id).order_by(models.ImportRow.row_number)))
    wb = Workbook()
    ws = wb.active
    ws.title = "导入结果"
    headers = ["Excel行号", "物料类型", "原输入编码", "预览／最终系统物料编码", "历史物料号", "关键器件码", "名称", "规格型号", "属性", "单位", "状态", "错误原因", "相似候选", "处理结果"]
    ws.append(headers)
    for row in rows:
        payload = json.loads(row.payload_json)
        similarity = json.loads(row.similarity_json)
        ws.append([
            row.row_number, {'material':'原材料','semi_finished':'半成品','unit':'单元','machine':'整机'}.get(payload.get('item_type'),'原材料'),
            payload.get('code_input'), payload.get("code"), payload.get('historical_item_code'), payload.get('key_component_code'), payload.get("name"), payload.get("specification"),
            SOURCE_LABELS.get(payload.get("source_type"),''), payload.get("unit"), {'ready':'待提交','yellow':'相似提示','red':'高度相似','error':'错误','imported':'已导入','cancelled':'已取消'}.get(row.status,row.status), row.error_message,
            "\n".join(f"{x['item']['code']} {x['item']['name']} ({x['full_score']}%)" for x in similarity),
            {'create':'创建','skip':'跳过'}.get(row.decision,row.decision),
        ])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, width in enumerate([10,12,22,26,22,16,28,36,12,10,14,36,50,16], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    return workbook_bytes(wb)


def audit_workbook(events: list[models.AuditEvent]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "变更日志"
    ws.append(["ID", "时间", "操作人ID", "动作", "对象类型", "对象ID", "原因", "批次", "修改前", "修改后"])
    for event in events:
        ws.append([
            event.id, event.created_at.isoformat(), event.actor_id, event.action, event.entity_type,
            event.entity_id, event.reason, event.batch_key, event.before_json, event.after_json,
        ])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(ws.max_row, 1)}"
    return workbook_bytes(wb)


def compare_workbook(comparison: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM差异"
    items = comparison["items"]
    ws.append(["物料编码", "名称", *[f"{item['code']} {item['name']}" for item in items], "是否差异"])
    different_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in comparison["rows"]:
        values = [row["values"].get(str(item["id"])) for item in items]
        ws.append([row["item"]["code"], row["item"]["name"], *values, "是" if row["different"] else "否"])
        if row["different"]:
            for cell in ws[ws.max_row]:
                cell.fill = different_fill
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    return workbook_bytes(wb)
