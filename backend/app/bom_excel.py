"""Excel-only layout. Browser BOM trees continue to use their existing row format."""
from __future__ import annotations

from decimal import Decimal

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import models, services
from .path_bom import EffectiveBOM


def candidate_signature(row):
    return tuple(sorted((int(m['item_id']), Decimal(str(m['market_share']))) for m in row.get('members', [])))


def candidate_cells(row, enabled=True):
    if not enabled:
        return []
    members = sorted(row.get('members', []), key=lambda m: int(m['item_id']) != row['item']['id'])
    result = []
    for member in members:
        item = member['item']
        name = item['name']
        if len(name) > 6:
            name = name[:3] + '..' + name[-3:]
        share = format(Decimal(str(member['market_share'])).normalize(), 'f')
        result.append(f"{item['code']}({name} {share}%)")
    return result


def write_sheet(ws, title, headers, records, candidates):
    count = max((len(values) for values in candidates), default=0)
    width = len(headers) + count
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    # Two header rows allow one merged group title and one selectable header per candidate.
    for col, header in enumerate(headers, 1):
        ws.cell(2, col, header)
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        ws.column_dimensions[get_column_letter(col)].width = 12 if header in ('序号', '单位', '数量') else 30
    if count:
        first = len(headers) + 1
        ws.cell(2, first, '选配组')
        if count > 1:
            ws.merge_cells(start_row=2, start_column=first, end_row=2, end_column=width)
        for i in range(count):
            ws.cell(3, first + i, '当前选用' if i == 0 else f'备选 {i}')
            ws.column_dimensions[get_column_letter(first + i)].width = 42
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=width):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    for row_index, (values, options) in enumerate(zip(records, candidates, strict=True), 4):
        for col_index, value in enumerate([*values, *options, *([None] * (count - len(options)))], 1):
            ws.cell(row_index, col_index, value)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            # Names/remarks from business input must remain text, not executable formulas.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(width)}{max(3, ws.max_row)}'


def export_workbook(db, root_id, *, production=False, show_alternatives=True, expand_materials=None):
    from .excel_io import latest_item_change_reasons, safe_filename, technical_level_labels, workbook_bytes
    root = db.get(models.Item, root_id)
    if not root:
        raise HTTPException(404, '对象不存在')
    expand = services.default_expand_materials(root) if expand_materials is None else expand_materials
    graph = EffectiveBOM(db)
    # Never traverse backup branches for an Excel export.
    actual = graph.rows(root_id, show_alternatives=False, expand_materials=expand)
    if not actual and root.item_type == 'material':
        actual = [{'sequence': 1, 'level': 0, 'path': [root.id], 'item': services.item_dict(root),
                   'quantity': '1', 'total_quantity': '1', 'members': [], 'is_terminal_material': True}]
    upper = [r for r in actual if r.get('members') and not r.get('is_terminal_material')]
    if production:
        merged = {}
        for row in actual:
            if not row.get('is_terminal_material'):
                continue
            key = (row['item']['id'], candidate_signature(row))
            target = merged.setdefault(key, {**row, 'sum': Decimal(0)})
            target['sum'] += Decimal(row['total_quantity'])
        rows = [{**r, 'quantity': r['sum']} for r in merged.values()]
    else:
        rows = actual
    headers = ['序号'] + ([] if production else ['层级']) + [
        '物料编码', '名称', '规格型号', '单位', '数量', '备注',
        '上一版本名称', '发票名称', '材料属性', '变更记录']
    reasons = latest_item_change_reasons(db, {r['item']['id'] for r in rows})
    levels = technical_level_labels(rows) if not production else []
    records = []
    for index, row in enumerate(rows):
        item = row['item']
        records.append([index + 1, *([] if production else [levels[index]]), item['code'], item['name'],
                        item.get('specification'), item['unit'], Decimal(row['quantity']),
                        (row.get('line_remark') if not production else None) or item.get('remark'),
                        item.get('previous_version_name'), item.get('invoice_name'), item.get('material_attribute'),
                        reasons.get(item['id'])])
    kind = '生产BOM' if production else '技术BOM'
    title = f"{root.code}-{root.name}｜显示选配：{'是' if show_alternatives else '否'}｜展开有下级组成的原材料：{'是' if expand else '否'}"
    if root.unofficial_status == 'archived':
        title += '｜封存来源，仅供历史参考'
    wb = Workbook()
    wb.active.title = kind
    write_sheet(wb.active, title, headers, records, [candidate_cells(r, show_alternatives) for r in rows])
    if not production:
        wb.active.column_dimensions['B'].width = 32
        for cells in wb.active.iter_rows(min_row=4, min_col=2, max_col=2):
            cells[0].alignment = Alignment(horizontal='left', vertical='top', wrap_text=False)
    if production and show_alternatives and upper:
        ws = wb.create_sheet('上层选配说明')
        write_sheet(ws, '上层选配说明（不参与原材料数量汇总）', ['序号', '组成路径', '物料编码', '名称'],
                    [[i + 1, ' → '.join(graph.items[v].code for v in r['path']), r['item']['code'], r['item']['name']]
                     for i, r in enumerate(upper)], [candidate_cells(r) for r in upper])
    return safe_filename(f'{root.code}-{root.name}-{kind}.xlsx'), workbook_bytes(wb)
