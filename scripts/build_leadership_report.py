"""Render the leadership report's original business diagrams and Word source."""
from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'docs/reports'
ASSETS = REPORT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
BLUE, TEAL, DARK, GRAY = '#2454A6', '#138A86', '#18314B', '#52677C'

def font(size, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc' if bold else 'C:/Windows/Fonts/msyh.ttc', size)

def canvas(height, title, subtitle):
    im = Image.new('RGB', (1600, height), '#F3F7FC')
    d = ImageDraw.Draw(im)
    d.text((65, 35), title, font=font(44, True), fill=DARK)
    d.text((65, 100), subtitle, font=font(26), fill=GRAY)
    return im, d

def box(d, rect, title, lines, color=BLUE):
    x, y, xx, yy = rect
    d.rounded_rectangle(rect, radius=22, fill='white', outline='#D6E2F0', width=3)
    d.rounded_rectangle((x+18,y+23,x+26,yy-23), radius=4, fill=color)
    d.text((x+48, y+30), title, font=font(31,True), fill=color)
    for i, line in enumerate(lines):
        d.text((x+48,y+86+i*42), line, font=font(25), fill=DARK)

def arrow(d, x, y, xx, yy, color=BLUE):
    d.line((x,y,xx,yy), fill=color, width=5)
    if yy == y:
        d.polygon([(xx,yy),(xx-14,yy-10),(xx-14,yy+10)], fill=color)
    else:
        d.polygon([(xx,yy),(xx-10,yy-14),(xx+10,yy-14)], fill=color)

im,d=canvas(520,'一套数据，贯通维护与输出','原材料、半成品、单元、整机统一管理')
cards=[('统一建档',['编码查重','名称、规格统一']),('维护组成',['层级展开','组合与数量维护']),('按需选配',['共用基础组成','按引用位置独立选']),('输出与追溯',['技术 / 生产清单','批量导出、反向查找'])]
for i,(title,lines) in enumerate(cards):
    x=55+i*385
    box(d,(x,185,x+335,375),title,lines,TEAL if i==2 else BLUE)
    if i<3: arrow(d,x+340,280,x+378,280)
d.rounded_rectangle((55,415,1545,480),radius=15,fill='#E3EEF8')
d.text((100,428),'全流程保障：变更留痕  ·  非法引用检查  ·  权限控制  ·  备份恢复',font=font(28),fill=DARK)
im.save(ASSETS/'bom-flow.png')

im,d=canvas(720,'共用一套组成，各机型独立选配','实际数量只算当前选用，备选与市场占比不参与用量计算')
box(d,(430,165,1170,320),'共用半成品 S',['某组成行数量：3　　候选组件：A / B'])
d.line((800,320,800,350),fill=BLUE,width=5)
d.line((400,350,1200,350),fill=BLUE,width=5)
arrow(d,400,350,400,385)
arrow(d,1200,350,1200,385,TEAL)
box(d,(65,395,775,620),'整机 M1：选用 A',['使用 S 数量 2 → A 用量 = 2 × 3 = 6','B 仅作备选，不计入用量'])
box(d,(825,395,1535,620),'整机 M2：选用 B',['使用 S 数量 2 → B 用量 = 2 × 3 = 6','A 仅作备选，不计入用量'],TEAL)
d.text((270,658),'M1 调整选配，不会误改 M2；共用数据与个性配置可以同时保留。',font=font(27),fill=DARK)
im.save(ASSETS/'bom-options.png')

doc=Document()
section=doc.sections[0]
section.top_margin=Inches(.65); section.bottom_margin=Inches(.65)
section.left_margin=Inches(.7); section.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Title'].font.size=Pt(25)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.color.rgb=RGBColor.from_string('2454A6')
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
source=REPORT/'BOM系统建设汇报-20260924.md'
text=source.read_text(encoding='utf-8')
supplement=REPORT/'BOM系统建设汇报-需求到实现补充.md'
if supplement.exists():
    text+='\n\n'+supplement.read_text(encoding='utf-8')
lines=text.splitlines()
i=0
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append(lines[i]); i+=1
        rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
        table=doc.add_table(rows=0, cols=3); table.style='Light Shading Accent 1'
        for row in rows:
            for cell,value in zip(table.add_row().cells,row.strip('|').split('|')):
                cell.text=value.strip()
        continue
    if line.startswith('!['):
        image_path=re.search(r'\]\((.+)\)',line).group(1)
        doc.add_picture(str(REPORT/image_path),width=Inches(6.7))
    elif line.startswith('# '): doc.add_heading(line[2:],0)
    elif line.startswith('## '): doc.add_heading(line[3:],1)
    elif line.startswith('### '): doc.add_heading(line[4:],2)
    elif line.startswith('#### '): doc.add_heading(line[5:],3)
    else: doc.add_paragraph(line)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('BOM 系统建设汇报 · 2026年9月24日')
output=REPORT/'BOM系统建设汇报-20260924.docx'
doc.save(output)
print('Created report with 2 original diagrams:', output)
